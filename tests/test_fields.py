"""「我的資料」攤平成欄位、以及哪些欄位會列給模型挑。

最重要的一件事在最下面：清單多一項就會掉格，所以少見的欄位一律「表格提到才列出來」。
"""
from __future__ import annotations

import tempfile
from pathlib import Path

import pytest
from docx import Document

from backend.core import filler

# 把新欄位都填滿的一份虛構資料——沒填值的欄位本來就不會列出來，
# 拿沒填的去測「有問就給」等於什麼都沒測
FULL = {
    "basic": {"name_zh": "虛構甲", "gender": "男", "children": "2",
              "birthday": "1996年04月15日",
              "military_branch": "陸軍", "military_rank": "下士",
              "military_start": "105年09月01日", "military_end": "106年08月31日"},
    "contact": {"mobile": "0900-000-000", "address_mailing": "虛構市虛構路 1 號",
                "postal_mailing": "100", "postal_household": "200"},
    "certificate": [{"name": "虛構證照", "issuer": "虛構機構",
                     "issued": "110年01月01日", "expires": "115年01月01日"}],
    "language": [{"name": "英文", "level": "中等", "listening": "中等",
                  "speaking": "中等", "reading": "精通", "writing": "中等"}],
    "preference": {"job_type": "全職", "industry": "資訊服務", "role": "後端工程師",
                   "shift": "可配合", "travel": "可短期出差"},
    "qa": {"strengths": "做事仔細", "weaknesses": "容易鑽牛角尖",
           "career_plan": "三年內成為資深工程師", "why_apply": "想做有規模的系統"},
    "experience": [{"company": "虛構科技", "title": "工程師",
                    "start": "110年01月01日", "end": "至今"}],
}

# 少見的欄位：表格沒問到就不該列給模型挑
GATED = [
    "contact.postal_mailing", "contact.postal_household",
    "basic.military_branch", "basic.military_rank",
    "basic.military_start", "basic.military_end", "basic.military_period",
    "basic.children", "basic.total_tenure",
    "certificate[0].issued", "certificate[0].expires",
    "language[0].name", "preference.job_type", "qa.strengths",
]


def form_with(labels: list[str]):
    """做一張只印著那些標籤的兩欄表格。"""
    doc = Document()
    table = doc.add_table(rows=len(labels), cols=2)
    for i, label in enumerate(labels):
        table.cell(i, 0).text = label
    path = Path(tempfile.mkdtemp()) / "t.docx"
    doc.save(path)
    _doc, form, _slots = filler.parse(path)
    return form


class TestDerived:
    """算出來的欄位：存著的值不算數，一律重算。"""

    def test_age_from_birthday(self):
        fields = filler.fields_of({"basic": {"birthday": "1996年04月15日"}})
        assert fields["basic.age"].isdigit()
        assert int(fields["basic.age"]) >= 29      # 存著的年齡去年填今年就錯了

    def test_stored_age_is_ignored(self):
        fields = filler.fields_of({"basic": {"birthday": "1996年04月15日", "age": "18"}})
        assert fields["basic.age"] != "18"

    def test_military_period_from_start_end(self):
        fields = filler.fields_of(FULL)
        assert "~" in fields["basic.military_period"]

    def test_total_tenure_follows_experience(self):
        one = filler.fields_of({"experience": [
            {"start": "2020年01月01日", "end": "2021年12月31日"}]})
        two = filler.fields_of({"experience": [
            {"start": "2018年01月01日", "end": "2021年12月31日"}]})
        assert one["basic.total_tenure"] != two["basic.total_tenure"]

    def test_unreadable_dates_are_not_guessed(self):
        fields = filler.fields_of({"basic": {"birthday": "不知道"}})
        assert "basic.age" not in fields

    @pytest.mark.parametrize("start,end", [
        ("2018年", "2020年"),            # 只有年份：DATE_RE 要年＋月
        ("2018", "2020"),
        ("民國107年", "民國109年"),
        ("不知道", "還在職"),
    ])
    def test_year_only_dates_do_not_crash(self, start, end):
        """使用者很自然會打「2018年」。以前 _months 回空字串，divmod 當場拋 TypeError，
        而 fields_of 在分析與匯出都會跑——等於每一次都當掉。"""
        fields = filler.fields_of({"experience": [
            {"company": "虛構公司", "title": "虛構職稱", "start": start, "end": end}]})
        assert "experience[0].tenure" not in fields    # 算不出來就不寫，不猜
        assert "basic.total_tenure" not in fields

    def test_total_tenure_skips_rather_than_undercounts(self):
        """一筆算得出來、一筆算不出來時，總年資要空著而不是只算得出來的那一筆。

        少算的年資會直接寫進履歷而且沒有任何提示，比空著更糟。
        """
        both = filler.fields_of({"experience": [
            {"start": "2018年01月01日", "end": "2019年12月31日"},
            {"start": "2020年", "end": "2021年"}]})
        assert "basic.total_tenure" not in both

    def test_row_without_dates_does_not_block_total(self):
        """整筆沒寫起訖的不算數：那是沒有主張期間，不是算不出來。"""
        fields = filler.fields_of({"experience": [
            {"start": "2018年01月01日", "end": "2019年12月31日"},
            {"company": "還沒填日期的那一筆"}]})
        assert fields["basic.total_tenure"] == "2年"


class TestGating:
    """表格提到才列給模型挑。"""

    @pytest.mark.parametrize("key", GATED)
    def test_not_offered_when_form_never_asks(self, key):
        plain = filler.usable_fields(form_with(["姓名", "行動電話", "地址"]), FULL)
        assert key not in plain

    @pytest.mark.parametrize("labels,expected", [
        (["通訊地址", "郵遞區號"], ["contact.postal_mailing", "contact.postal_household"]),
        (["證照名稱", "發照日期"], ["certificate[0].issued"]),
        (["證照名稱", "有效期限"], ["certificate[0].expires"]),
        (["軍種", "軍階", "入伍日期", "退伍日期"],
         ["basic.military_branch", "basic.military_rank",
          "basic.military_start", "basic.military_end"]),
        (["服役期間"], ["basic.military_period"]),
        (["子女人數"], ["basic.children"]),
        (["總年資"], ["basic.total_tenure"]),
        (["語文能力", "聽", "說", "讀", "寫"], ["language[0].name", "language[0].listening"]),
        (["期望工作型態", "可否輪班"], ["preference.job_type", "preference.shift"]),
        (["您的優點", "生涯規劃"], ["qa.strengths", "qa.career_plan"]),
    ])
    def test_offered_when_form_asks(self, labels, expected):
        fields = filler.usable_fields(form_with(labels), FULL)
        for key in expected:
            assert key in fields

    def test_asking_half_offers_half(self):
        fields = filler.usable_fields(form_with(["發照日期"]), FULL)
        assert "certificate[0].issued" in fields
        assert "certificate[0].expires" not in fields

    def test_other_peoples_data_still_gated(self):
        """家人、諮詢人本來就是這樣擋的，不能因為改寫而壞掉。"""
        profile = {**FULL, "family": [{"name": "虛構乙", "relation": "父"}]}
        assert "family[0].name" not in filler.usable_fields(form_with(["姓名"]), profile)
        assert "family[0].name" in filler.usable_fields(form_with(["家庭成員", "稱謂"]), profile)


class TestRowChoices:
    """一列一筆的表：另外一條規則——表上有問、自己也真的填了，才列出來。"""

    def test_filled_is_offered(self):
        assert filler._worth_offering(
            "certificate[].issued", {"certificate[0].issued": "110年01月01日"})

    def test_blank_is_not_offered(self):
        """履歷表印著「發照日期」但沒填，列出來只會佔掉一欄。"""
        assert not filler._worth_offering("certificate[].issued", {"certificate[0].name": "證照"})

    def test_ungated_fields_are_always_offered(self):
        assert filler._worth_offering("family[].contact", {})
        assert filler._worth_offering("experience[].salary", {})
