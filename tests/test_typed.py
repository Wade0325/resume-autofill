"""對映清單裡直接把某一格改成自己打的字。

手打的值比任何判斷都優先，而且照原樣寫——從「我的資料」推出該寫什麼的那些規則
（日期拆進年月日、單位前只收數字）不該再套一次，人家就是要那幾個字。
"""
from __future__ import annotations

import pytest
from docx import Document

from backend import db, service
from backend.core import planner


@pytest.fixture
def job_with_name(make_job):
    """一份工作：第一個空格判給中文姓名，第二個空格沒判到。"""
    blanks = make_job.blanks
    job_id = make_job(decided={blanks[0].id: ["basic.name_zh", 0, "model", "中文姓名"]})
    return job_id, blanks[0], blanks[1]


def row_of(payload, slot_id):
    return next(i for i in payload["items"] if i["slot_id"] == slot_id)


class TestTypedValue:
    def test_starts_as_the_model_guess(self, client, job_with_name):
        job_id, named, _blank = job_with_name
        row = row_of(client.get(f"/api/jobs/{job_id}").json()["plan"], named.id)
        assert (row["value"], row["source"]) == ("虛構甲", "model")

    def test_typing_overrides_the_model(self, client, job_with_name, text_of):
        job_id, named, _blank = job_with_name
        r = client.patch(f"/api/jobs/{job_id}/value",
                         json={"slot_id": named.id, "value": "王大明（手打）"})
        row = row_of(r.json(), named.id)
        assert (row["value"], row["source"], row["status"]) == ("王大明（手打）", "typed", "fill")

        client.post(f"/api/jobs/{job_id}/output", json={})
        out = text_of(service.output_path(job_id))
        assert "王大明（手打）" in out
        assert "虛構甲" not in out

    def test_can_type_into_a_cell_nobody_claimed(self, client, job_with_name, text_of):
        job_id, _named, blank = job_with_name
        r = client.patch(f"/api/jobs/{job_id}/value",
                         json={"slot_id": blank.id, "value": "補一句"})
        row = row_of(r.json(), blank.id)
        assert (row["value"], row["status"]) == ("補一句", "fill")

        client.post(f"/api/jobs/{job_id}/output", json={})
        assert "補一句" in text_of(service.output_path(job_id))

    def test_clearing_returns_to_the_model_guess(self, client, job_with_name):
        job_id, named, _blank = job_with_name
        client.patch(f"/api/jobs/{job_id}/value", json={"slot_id": named.id, "value": "先改掉"})
        r = client.patch(f"/api/jobs/{job_id}/value", json={"slot_id": named.id, "value": ""})
        row = row_of(r.json(), named.id)
        assert (row["value"], row["source"]) == ("虛構甲", "model")

    def test_only_spaces_counts_as_clearing(self, client, job_with_name, text_of):
        job_id, _named, blank = job_with_name
        client.patch(f"/api/jobs/{job_id}/value", json={"slot_id": blank.id, "value": "補一句"})
        r = client.patch(f"/api/jobs/{job_id}/value", json={"slot_id": blank.id, "value": "   "})
        row = row_of(r.json(), blank.id)
        assert (row["value"], row["status"]) == ("", "skip")

        client.post(f"/api/jobs/{job_id}/output", json={})
        assert "補一句" not in text_of(service.output_path(job_id))

    def test_profile_is_never_touched(self, client, job_with_name):
        job_id, named, _blank = job_with_name
        before = db.get_kv("profile")
        client.patch(f"/api/jobs/{job_id}/value", json={"slot_id": named.id, "value": "手打的"})
        assert db.get_kv("profile") == before

    def test_typed_values_do_not_follow_to_another_job(self):
        assert service.typed_values({"typed": None}) == {}


class TestTypedWins:
    def test_typed_beats_the_apply_panel(self, client, make_job):
        """「這次應徵」也是使用者填的，但手打那一格更直接，以它為準。"""
        title_slot = next((s for s in make_job.slots if s.job_field == "job.title"), None)
        if title_slot is None:
            pytest.skip("這份樣本表格沒有應徵職務欄")
        job_id = make_job()
        client.patch(f"/api/jobs/{job_id}/apply", json={"values": {"job.title": "面板填的職務"}})
        r = client.patch(f"/api/jobs/{job_id}/value",
                         json={"slot_id": title_slot.id, "value": "手打的職務"})
        row = row_of(r.json(), title_slot.id)
        assert (row["value"], row["source"]) == ("手打的職務", "typed")

    def test_unit_cells_keep_what_you_typed(self, client, make_job, text_of, tmp_path):
        """日期那種「數字＋單位」的格子，自動判斷只收數字，手打的照寫。

        樣本表格沒有這種版面，所以這裡自己做一份「＿年＿月＿日」的。
        """
        from backend.core import filler

        doc = Document()
        t = doc.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "出生日期"
        t.cell(0, 1).text = "　　　年　　　月　　　日"
        form = tmp_path / "with_units.docx"
        doc.save(str(form))

        gap = next((s for s in filler.parse(form)[2]
                    if s.kind == "gap" and s.option in ("年", "月", "日")), None)
        assert gap is not None, "這份表格應該要有帶單位的空格"

        job_id = make_job(form=form, decided={})
        client.patch(f"/api/jobs/{job_id}/value",
                     json={"slot_id": gap.id, "value": "民國九十九"})
        client.post(f"/api/jobs/{job_id}/output", json={})
        assert "民國九十九" in text_of(service.output_path(job_id))


class TestClassicRoute:
    """讀文字那條路也要吃手打的值。"""

    def test_fills_a_cell_nobody_claimed(self, sample_form):
        from backend.core import document

        _text, slots = document.ParsedDoc(str(sample_form)).flatten()
        target = slots[0]
        ops, _skipped = planner.build_plan(slots, {}, {}, {target.id: "讀文字手打"})
        assert [(o.slot.id, o.value, o.source) for o in ops] == [
            (target.id, "讀文字手打", "typed")]

    def test_beats_the_model(self, sample_form):
        from backend.core import document

        _text, slots = document.ParsedDoc(str(sample_form)).flatten()
        target = slots[0]
        decided = {target.id: planner.Decision("basic.name_zh", 0, "model", "中文姓名")}
        ops, _ = planner.build_plan(slots, {"basic": {"name_zh": "虛構甲"}}, decided,
                                    {target.id: "蓋過模型"})
        assert [(o.value, o.source) for o in ops] == [("蓋過模型", "typed")]

    def test_without_typing_the_model_stands(self, sample_form):
        from backend.core import document

        _text, slots = document.ParsedDoc(str(sample_form)).flatten()
        target = slots[0]
        decided = {target.id: planner.Decision("basic.name_zh", 0, "model", "中文姓名")}
        ops, _ = planner.build_plan(slots, {"basic": {"name_zh": "虛構甲"}}, decided)
        assert [(o.value, o.source) for o in ops] == [("虛構甲", "model")]
