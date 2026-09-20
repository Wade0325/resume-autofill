"""測試共用的準備工作。

兩件事必須在 import backend 之前做完，所以寫在模組最上面、不放進 fixture：

1. `RESUME_AUTOFILL_HOME` 指到暫存資料夾——`backend.config` 是在 import 當下讀環境變數的，
   晚一步設就會寫到真正的 `data/`，把使用者的履歷蓋掉。
2. 推論服務指到一個沒人在聽的埠。整套測試都不需要模型：要模型判斷的地方一律換成假的，
   要測「模型沒開」的行為時就讓它真的連不上。
"""
from __future__ import annotations

import os
import shutil
import sys
import tempfile
import uuid
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

_HOME = Path(tempfile.mkdtemp(prefix="resume_autofill_tests_"))
os.environ["RESUME_AUTOFILL_HOME"] = str(_HOME)
os.environ["RESUME_AUTOFILL_LLM_HOST"] = "http://127.0.0.1:8099"   # 故意沒人在聽
os.environ["RESUME_AUTOFILL_AUTOSTART"] = "0"                      # 別去啟動模型

from backend import config, db, service  # noqa: E402
from backend.core import filler  # noqa: E402

config.ensure_dirs()
db.init()

# 整份測試用的虛構資料。不用任何真實履歷——研究用的考題是本人的真實資料，
# 沒進版控，測試不能依賴它
PROFILE = {
    "basic": {"name_zh": "虛構甲", "name_en": "Chia Hsu-Kou", "gender": "男",
              "birthday": "1996年04月15日", "military": "役畢",
              "marital_status": "未婚", "national_id": "A123456789"},
    "contact": {"mobile": "0900-000-000", "email": "test@example.invalid",
                "address_household": "虛構市虛構區虛構路 1 號",
                "address_mailing": "虛構市虛構區虛構路 1 號"},
    "education": [
        {"school": "虛構大學", "department": "資訊工程學系", "degree": "學士",
         "status": "畢業", "start": "2014年09月01日", "end": "2018年06月30日"},
    ],
    "experience": [
        {"company": "虛構科技", "title": "後端工程師", "start": "2018年08月01日",
         "end": "至今", "salary": "50,000", "leave_reason": "在職中",
         "description": "負責後端服務的設計與維護"},
    ],
    "job": {"expected_salary": "面議", "available_date": "隨時"},
    "autobiography": "我是虛構甲，這段自傳純屬測試用。",
}


@pytest.fixture(scope="session")
def home() -> Path:
    """這一輪測試的資料夾。真正的 data/ 不會被碰到。"""
    return _HOME


@pytest.fixture(scope="session")
def sample_form(tmp_path_factory) -> Path:
    """一份虛構的空白履歷表：`tools/make_sample.py` 產的，格式仿台灣公司的表格。

    有標籤在左的基本資料、勾選框、一列一筆的學經歷表、整格的自傳，
    還有「以下由人事單位填寫」——填寫端要處理的狀況大多在這一份裡面。
    """
    sys.path.insert(0, str(ROOT / "tools"))
    import make_sample

    path = tmp_path_factory.mktemp("form") / "sample_resume_form.docx"
    make_sample.build(str(path))
    return path


@pytest.fixture(autouse=True)
def clean_state():
    """每一項測試都從乾淨的共用狀態開始。

    HOME 是整輪共用的（config 在 import 時就決定了，改不了），所以改成每次把
    共用的那幾張表清掉；工作各自用隨機代碼，不會互相干擾。
    """
    with db.connect() as conn:
        conn.execute("DELETE FROM kv")
        conn.execute("DELETE FROM template")
        conn.execute("DELETE FROM job")
        conn.execute("DELETE FROM import_job")
        conn.execute("DELETE FROM profile_version")
    db.put_kv("profile", PROFILE)
    yield


@pytest.fixture
def client():
    """打 API 用。用 with 進去才會跑 startup（它會把殘留的工作標成中斷）。"""
    from fastapi.testclient import TestClient

    from backend.main import app

    with TestClient(app, base_url="http://127.0.0.1") as c:
        yield c


@pytest.fixture
def make_job(sample_form):
    """生一份「已經分析完」的工作，不必真的跑模型。

    decided 給 {位置編號: [欄位代碼, 第幾筆, 來源, 標籤]}；不給就挑第一個空格填中文姓名。
    """
    _doc, _form, slots = filler.parse(sample_form)
    anchors = [{"id": s.id, "kind": s.kind, "label": "", "tick": None, "basis": None}
               for s in slots]
    blanks = [s for s in slots if s.kind == "blank"]

    def _make(filename: str = "虛構公司表格.docx", *, status: str = "analyzed",
              decided=None, engine: str = "vlm", stage: str = "", error: str = "",
              form=None):
        """form 給別份表格就用那一份（要測特殊版面時用），不給就用共用的樣本。"""
        source = form or sample_form
        rows = anchors if form is None else [
            {"id": s.id, "kind": s.kind, "label": "", "tick": None, "basis": None}
            for s in filler.parse(source)[2]]
        job_id = uuid.uuid4().hex[:12]
        service.job_dir(job_id).mkdir(parents=True, exist_ok=True)
        shutil.copy(source, service.input_path(job_id))
        db.create_job(job_id, filename, status="processing", engine=engine)
        if status == "analyzed":
            db.update_job(job_id, status="analyzed", fingerprint=f"{engine}:test",
                          decided=decided if decided is not None else
                          {blanks[0].id: ["basic.name_zh", 0, "model", "中文姓名"]},
                          anchors=rows)
        elif status == "failed":
            db.update_job(job_id, status="failed", error=error or "分析失敗了")
        elif stage:
            db.update_job(job_id, stage=stage)
        return job_id

    _make.slots = slots
    _make.blanks = blanks
    return _make


@pytest.fixture
def text_of():
    """把 .docx 的所有文字拉成一段，用來確認值有沒有寫進去。"""
    def _text(path_or_bytes) -> str:
        import io

        import docx
        src = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else str(path_or_bytes)
        d = docx.Document(src)
        parts = [p.text for p in d.paragraphs]
        parts += [c.text for t in d.tables for row in t.rows for c in row.cells]
        return "\n".join(parts)
    return _text
