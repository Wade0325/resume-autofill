# -*- coding: utf-8 -*-
r"""空白履歷表 ＋ 個人資料（app.db 那份 JSON）→ 填好的履歷表。

視覺版的填寫引擎：讓模型看著版面示意圖決定「每一個位置放哪一項資料」。
另一條路是 planner.py（規則錨定＋純文字模型），兩條都在，由 service 決定用哪一條。

這個檔案同時是研究對象：claude_code_in_agent/ 那套評分與研究迴圈跑的就是它，
改進填寫效果就是改這裡。介面只有三個：analyze()、write()、solve()。

作法：程式決定「哪裡能寫、怎麼寫」，模型只決定「這個位置放哪一項資料」。

  1. 照 docx 的表格骨架自己畫版面示意圖，每格標上地址 —— 讓 VLM 看得到位置
  2. 機械地列出所有「可以寫字的位置」：空格子、印好的字之間的留白、冒號後面、
     勾選框。「以下由公司填寫」之後、應徵職務這種每間公司不同的欄位不列
  3. 一列一筆的表（學經歷、家人）整張問模型「每一欄是什麼」，列的順序由程式排：
     列首印著學位的照學位放、有序號欄的照序號、都沒有才由上而下
  4. 其餘位置分批問模型：這裡該填哪一項資料（挑不到就 __無__）。答案被 JSON Schema
     約束成只能挑既有的項目代碼，模型編不出不存在的資料
  5. 補漏：還空著的位置，只拿名稱相近、還沒用到的幾項資料再問一次
  6. 過濾與去重：放錯地方的配對擋掉（詞對不上、表格沒提到的別人的資料、選項值
     離開了題目），同一項資料在同一張表只留欄名最像的那一格
  7. 字面對不上的勾選題（資料「無」、表格印「□否」）交給模型判斷意思；
     是非題只讓模型判斷相不相關，勾哪個由程式決定
  8. 程式把值寫進去：日期照單位拆格、西元換民國、期間由起訖重算。
     原本印在表格上的字一個都不會動

第一版是讓模型輸出「整格的新內容」，三種錯誤同時發生：值被寫進標題格、
憑空捏造沒有的資料、順手改掉原本印的標點（40 格只對 13 格，另外多填 68 格）。
位置與寫法收回程式手上之後，這三類錯誤從結構上消失。

實測結果與每一項修正的效果見同目錄的 notes.md。

單獨試跑一份與逐格評分：claude_code_in_agent/README.md
"""
from __future__ import annotations

import base64
import io
import logging
import os
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from datetime import date
from itertools import groupby, permutations
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from PIL import Image, ImageDraw, ImageFont

from . import llm
from .document import ROC_BEFORE_RE, ROC_WORD_RE, iter_block_items, to_roc
from .runs import write_changes
from .schema import (
    BY_KEY,
    LABEL_ALIASES,
    OPTION_SYNONYMS,
    PER_JOB_LABELS,
    PRESENT_RE,
    PRESENT_WORDS,
)

log = logging.getLogger(__name__)

# 預設值只給研究用的入口；產品一律由 service 把 config 裡的設定傳進來。
# 寫 IP 不寫 localhost 的原因見 config.LLM_HOST
LLM_HOST = "http://127.0.0.1:8085"
LLM_MODEL = "Qwen3.5-9B-Q4_K_M"

# 版面示意圖：自己畫，不借助任何外部排版引擎（LibreOffice、Word 都不能用——
# 產品是可攜、離線的單機工具，裝不進去的東西就不能進解法）
FONT_PATH = os.environ.get("FORM_FONT", r"C:\Windows\Fonts\msjh.ttc")
PAGE_W, PAGE_H, MAX_PAGES = 1100, 1500, 4
FONT_SIZE, ADDR_SIZE, LINE_H, ADDR_H = 15, 11, 20, 14

# 使用者在對映清單自己打的值：接在同一套寫入流程上，代碼加前綴跟欄位代碼區分開
TYPED = "__typed__"

BATCH = 16     # 一次問幾個位置。問多了模型會整批放棄（見 backend/core/planner.py 的實測）
BOX_BATCH = 8  # 一次問幾題勾選題。二十題一起問，宣告事項那七題會整批答不出來
NONE = "__無__"

CHECKBOX_CHARS = "□☐▢◻"
CHECKED = "■"
# 沒有印框、而是「欄名是選項、底下空一格」的表（區部：日間｜夜間，狀態：畢業｜肄業）
# 在格子裡打的記號
MARK = "✓"
# 印好的字之間留出來的書寫空間：「自　　年　　月」「血型：　　型」。
# 兩格以上的空白才算，全形半形混著數；剛好兩格的、以及行首的，要後面接著單位才算
# ——「年  月  日」的月日前面只有兩格，而「cm  體重」中間那兩格只是排版間隔，
# 「　　　　初試日期：」開頭那一長串則是縮排。
GAP_RE = re.compile(r"[ 　]{2,}|[_＿]{2,}|[.．]{4,}")
# 印在空格後面的單位。值一律寫在單位前面（「＿＿年」「＿＿公分」），認得單位有兩個
# 用處：判斷一段空白是不是留白，以及日期要拆成哪幾格
UNITS = ("公分", "公斤", "cm", "kg", "年", "月", "日", "時", "分", "秒", "歲", "型", "元")
DATE_UNITS = ("年", "月", "日", "時", "分", "秒")
# 單位後面可以再接的字：「＿月＿日後」的「日」一樣是單位
UNIT_SUFFIX = "後前起"
# 台灣的市話區碼，長的排前面（「(　　)＿＿」要把 0224596466 拆成 02 與 24596466）
AREA_CODES = ("0836", "0826", "089", "082", "049", "037",
              "02", "03", "04", "05", "06", "07", "08")
DATE_RE = re.compile(r"(\d{2,4})\s*[年/.\-]\s*(\d{1,2})(?:\s*[月/.\-]\s*(\d{1,2}))?")
PLACEHOLDER_RE = re.compile(r"^[\s　_＿…．.\-—–]*$")
# 多選值的分隔符號。刻意不含空白：「2026 年 8 月 24日」用空白拆會拆出一個「月」，
# 剛好跟可到職日的選項「□　月　日」字面相同，就被當成已經勾好、日期不寫了
OPTION_SPLIT_RE = re.compile(r"[、，,／/；;]+")
# 「※以下欄位由本公司人員填寫※」——表格自己說了後面不是求職者填的
COMPANY_ONLY_RE = re.compile(r"以下.{0,8}(公司|人事|人資).{0,8}填")
# 公司區塊之後又回到應徵者的部分。要先比這個：「以下由應徵者填寫，公司人員勿填」兩個都對得上
APPLICANT_AGAIN_RE = re.compile(r"以下.{0,8}(應徵者|應徵人|求職者|申請人|本人).{0,8}填")
# 別人的資料：表格上沒印這些字，就不拿出來給模型挑（同產品匯入端 reader._SECTION_GATES）
OTHER_PEOPLE = {
    "emergency": ("緊急聯絡", "緊急連絡"),
    "family": ("家庭", "家屬", "家人", "父", "母"),
    "reference": ("推薦", "諮詢", "介紹人"),
}
# 表格提到才列給模型挑。別人的資料（上面那三區）本來就只在提到那個人時才拿出來；
# 語言、求職偏好、問答題多數表格根本沒有，清單越長模型越容易配錯
_ROW_INDEX = re.compile(r"\[\d+\]")

ASK_ONLY_IF_MENTIONED = {
    **OTHER_PEOPLE,
    "language": ("語文", "語言", "英文", "外語", "母語"),
    "preference": ("工作型態", "期望產業", "職務類別", "輪班", "外派", "出差", "求職條件"),
    "qa": ("優點", "缺點", "生涯", "規劃", "動機", "抱負", "期許"),
    # 單獨一個欄位也能這樣擋。多一個選項就多一次配錯的機會：實測把發照日期、
    # 郵遞區號無條件放進清單，模型在密集的證照表與聯絡欄就開始挑錯格子
    "contact.postal_mailing": ("郵遞區號", "郵區", "郵遞"),
    "contact.postal_household": ("郵遞區號", "郵區", "郵遞"),
    "certificate[].issued": ("發照", "發證", "取得日"),
    "certificate[].expires": ("到期", "有效期"),
    "basic.military_branch": ("軍種",),
    "basic.military_rank": ("軍階", "階級"),
    "basic.military_start": ("入伍",),
    "basic.military_end": ("退伍",),
    "basic.military_period": ("服役期間",),
    "basic.children": ("子女",),
    "basic.total_tenure": ("總年資", "總計年資"),
}
# 公司自己填的欄位。「以下由公司填寫」那條線之前也會夾雜這種欄位（面談日期印在最上面）。
# 「到職日期」是公司填的報到日；前面有「可」的「可到職日期」是應徵者填的，不能一起擋掉
COMPANY_WORDS_RE = re.compile(r"面談|初試|複試|任用|建議薪資|主管簽章|(?<!可)到職日期")
# 親筆簽名留給本人手寫；同一行的「日期＿年＿月＿日」是簽名的日期，個人資料也不會有
SIGN_WORDS = ("簽名", "簽章")
# 被擋的詞怎麼擋：簽名、公司欄位、「年制」（「□二年制 □四年制」整排都是選項）一出現，
# 整行都不是應徵者要填的；應徵職務、工作地點這種每間公司不一樣的欄位只算自己那一段，
# 同一行的「希望待遇：＿＿」照填。這些位置留著、標上它是哪一個欄位，值由填寫頁的
# 「這次應徵」面板提供（沒填就留白）；模型看不到它們，位置編號也另外算，
# 學過的格式才不會因此全部對不上
JOB_SPECIFIC = tuple(PER_JOB_LABELS)
# 兩邊都常出現、卻不代表相關的詞：履歷表的問答題幾乎都有「工作」；
# 「期間」則是服役期間、就學期間、任職期間都有，光靠它會把學歷填進服役欄
GENERIC_WORDS = {"工作", "期間"}
# 否定的選項或值：「□否」「□無」「不同意」
NEGATIVE = r"否|無|沒有|不.*"

SYSTEM = """你是履歷表格的填寫助理。表格上每一個「可以寫字的位置」都編了號，
位置在那一格的字裡用 ▁ 標出來。請為每一個位置指出：該填個人資料裡的哪一項。

規則：
1. 只能從個人資料的項目代碼裡挑，挑不到就填 __無__。值長什麼樣不用你管，
   程式會照個人資料原樣填進去。
2. 這些位置一律填 __無__：公司自己要填的（初試日期、複試日期、面談情形、
   任用與否、建議薪資、主管簽章）、個人資料裡沒有那項的、以及只是排版留白的地方。
3. ▁ 原本是 □ 的位置是勾選框：填「這個選項在問哪一項資料」就好。勾哪一個由程式
   比對資料的值跟選項決定，意思相同也算（資料「無」對選項「否」）。
   例如「□未婚 □已婚」兩個位置都填 basic.marital_status 是對的。
4. 多筆的資料（學歷、工作經歷、家人）看上面欄名後面標的「往下第幾列」：
   第 1 列填 [0]、第 2 列填 [1]，依此類推；資料沒有那一筆就填 __無__。
5. 表格問的若是別人的資料（家庭成員、緊急聯絡人、推薦人），個人資料裡沒有那個人，
   就填 __無__，不要拿本人的姓名生日頂替。
6. 上面兩條是為了不要亂填，但也不要因為不確定就一律填 __無__：
   左邊欄名或上面欄名對得上個人資料裡的某一項，那就是它，填下去。
7. 每一個位置都要回答一次。"""


@dataclass
class Cell:
    """文件裡一個可以寫字的容器（表格的一格，或內文的一段）。"""
    addr: str
    paras: List[Any]
    row_head: str = ""      # 左邊最近一格印的字（欄名）
    col_head: str = ""      # 同一欄上面印的字
    depth: int = 0          # 空格子在上面欄名底下第幾列（一列一筆的表才有）
    printed_in_row: int = 0    # 這一列有幾格印著欄位名稱（欄名列不會只有一兩格）
    value_after: bool = False  # 右邊還有空格子（代表這格是欄名，不是提示字）
    caption: str = ""          # 欄名列正上方那一句說明（「請列舉…並同意我們諮詢」）
    head_row: int = -1         # col_head 印在第幾列
    col_parent: str = ""       # 欄名上面那一層（「區部」底下分「日間｜夜間」）
    origin: int = -1           # 一列一筆的表：表頭從第幾列開始，同一張表的格子這個值相同
    key: str = ""              # 列首印的這一筆是哪一種（「碩／博士」「大學」），沒有就空
    seq: int = 0               # 左邊序號欄印的數字（「1」「2」），沒有就 0
    row_first: str = ""        # 整列都是字的那種列（「緊急聯絡人｜關係｜聯絡電話：…」）最左邊那格


@dataclass
class Slot:
    """一個可以寫字的位置。程式知道怎麼寫，模型只要說「這裡放哪一項資料」。"""
    id: str
    kind: str               # blank | gap | append | line | box
    addr: str
    para: int               # 在 Cell.paras 的第幾段
    start: int
    end: int
    filler: str = ""        # gap 原本的留白，用來保持欄寬
    option: str = ""        # box：選項印的字；gap：緊接在後面的單位（年、月、kg），有單位只收數字
    preview: str = ""       # 給模型看的那一格內容，目標位置標成 ▁
    # 「這次應徵」才有的欄位（應徵職務、工作地點）：值跟著那一份工作，不在「我的資料」裡
    job_field: str = ""     # 對到的欄位代碼，例如 job.title
    extra: bool = False     # 這個位置不給模型看、也不進格式指紋（見 JOB_SPECIFIC）
    cell: Cell = field(default=None, repr=False)


# ---------------------------------------------------------------------------
# 一、把文件拆成「格子」與「可以寫字的位置」
# ---------------------------------------------------------------------------

def _unit_at(text: str, pos: int) -> str:
    """text 從 pos 開始印的是哪一個單位，沒有就回空字串。

    後面還接著中英文就不算，不然「每月薪資」的「月」也會被當成單位。
    """
    rest = text[pos:]
    for unit in sorted(UNITS, key=len, reverse=True):
        after = rest[len(unit):].lstrip(UNIT_SUFFIX) if rest.startswith(unit) else ""
        if rest.startswith(unit) and not re.match(r"[一-鿿A-Za-z]", after):
            return unit
    return ""


_TEMPLATE_RE = re.compile(r"[\s　\d自至到~～\-—–－/／()（）]|"
                          + "|".join(sorted(UNITS, key=len, reverse=True)))


def _is_template(text: str) -> bool:
    """整格只印著格式，不是欄位名稱也不是值：「自　　年　　月／至　　年　　月」、
    「　　公分／　　公斤」。這種格子跟空格子一樣是拿來填的。"""
    return bool(text.strip()) and not _TEMPLATE_RE.sub("", text).strip()


def _box_only(text: str) -> bool:
    """整格只有勾選框跟很短的選項字：「□畢 □肄」。一列一筆的表也會整欄都是這種格子。"""
    body = re.sub(f"[{CHECKBOX_CHARS}{CHECKED}]", "", text)
    return (any(c in text for c in CHECKBOX_CHARS + CHECKED)
            and len(_squash(body)) <= 6 and not re.search(r"[：:]", text))


def _merged_away(tc) -> bool:
    """這一格是垂直合併的延續格——字都在最上面那格裡，寫進來看不見。"""
    pr = tc.tcPr
    vmerge = None if pr is None else pr.find(qn("w:vMerge"))
    return vmerge is not None and vmerge.get(qn("w:val"), "continue") == "continue"


def _head(heads: List[Tuple[int, int, list]], edges: List[int], index: int, count: int,
          col: int, span: int) -> Tuple[str, int, int, int]:
    """這一格頭上最近的欄名：(字, 在第幾列, 起始網格欄, 結束網格欄)，找不到列號是 -1。

    上面那一列跟這一列格數相同時，看同一個位置的那一格。網格欄（gridCol）常常
    對不齊：學歷表「在學」底下那一格的網格從「肄業」中間開始，照網格欄找會找到
    肄業，Word 畫出來卻在「在學」底下——一格一格對著排的表，第幾格才是真的。
    格數不同時（上面一格「身分證字號」，底下切成十格）才看寬度重疊得最多的那一格。
    """
    lo, hi = edges[col], edges[min(col + span, len(edges) - 1)]
    for row_no, above_count, labels in reversed(heads):
        if count is None:       # 只認網格起點相同的（不在表格清單裡的格子，沿用原本的找法）
            hit = next((x for x in labels if x[1] == col), None)
            if hit:
                return hit[3], row_no, hit[1], hit[2]
            continue
        if above_count == count:
            hit = next((x for x in labels if x[0] == index), None)
            if hit:
                return hit[3], row_no, hit[1], hit[2]
            continue
        last = len(edges) - 1
        overlap = [(min(hi, edges[min(c1, last)]) - max(lo, edges[min(c0, last)]), c0, c1, text)
                   for _i, c0, c1, text in labels]
        best = max(overlap, key=lambda x: (x[0], -x[1]), default=None)
        if best and best[0] > 0:
            return best[3], row_no, best[1], best[2]
    return "", -1, col, col + span


def cells(doc) -> List[Cell]:
    """列出所有格子。地址與 evaluate.py 的報告一致，對錯照著地址查得到。

    直接數 XML 裡的 tc（這一列實際存在的儲存格），不用 row.cells——合併儲存格
    在 row.cells 會重複出現，而 lxml 元素代理是隨用隨建的，拿 id() 去重會時多時少。
    順便記下同列左邊、同欄上面印的字：判斷「這一格是什麼」全靠它們。
    """
    out: List[Cell] = []
    table_no = para_no = 0
    # 「以下由公司填寫」之後是公司自己的欄位（面談情形、任用與否、建議薪資），不列。
    # 以前碰到這句就整份不看了、而且只認段落：寫在表格裡的沒認出來，公司區塊夾在中間時
    # 後面的應徵者部分也一起丟掉。現在表格裡的這句也認，看到「以下由應徵者填寫」就回來。
    # 跳過的段落、表格照樣算編號，地址才跟 evaluate.py 對得上
    company = False
    for block in iter_block_items(doc):
        if isinstance(block, Table):
            edges = [0]
            for g in block._tbl.tblGrid.gridCol_lst:
                edges.append(edges[-1] + int(g.w or 0))
            heads: List[Tuple[int, int, list]] = []   # (列, 這一列幾格, [(第幾格, 起, 迄, 字)])
            header_rows = set()                     # 整列都是欄位名稱的列（表頭）
            cap_row, cap_text = -2, ""          # 最近一列「整列只有一句長說明」的列
            for r, row in enumerate(block.rows):
                tcs = row._tr.tc_lst
                row_text = " ".join(_Cell(tc, block).text for tc in tcs)
                if APPLICANT_AGAIN_RE.search(row_text):
                    company = False
                    continue            # 這一列本身是說明，不是填寫位置
                if company or COMPANY_ONLY_RE.search(row_text):
                    company = True
                    continue
                grid_col, in_row = 0, []
                for c, tc in enumerate(tcs):
                    col, span = grid_col, tc.grid_span or 1
                    grid_col += span
                    head = _head(heads, edges, c, len(tcs), col, span)
                    if _merged_away(tc):
                        # 延續格不是填寫位置，但上面那格的字仍是它右邊格子的欄名：
                        # 「聯絡電話」跨兩列，第二列的「(H)：」左邊就是它
                        in_row.append((c, None, None, _head(heads, edges, c, None, col, span)[0],
                                       head, col, col + span))
                        continue
                    paras = _Cell(tc, block).paragraphs
                    text = "\n".join(p.text for p in paras).strip()
                    in_row.append((c, f"t{table_no}.r{r}.c{c}", paras, text, head,
                                   col, col + span))

                def is_label(t):
                    return bool(t) and not _is_template(t) and not _box_only(t)

                real = [x for x in in_row if x[1]]
                labels = [x for x in real if is_label(x[3])]
                # 整列只有一句長說明：那是這一區在講什麼。諮詢人那張表沒有左邊的
                # 區塊標題，欄名又只有「姓名／職稱／公司名稱」，跟家庭成員長得一樣，
                # 只有上面那句「請列舉…並同意我們諮詢」認得出它問的是誰
                if len(labels) == 1 and len(_squash(labels[0][3])) > 8:
                    cap_row, cap_text = r, labels[0][3]
                # 列首印著這一筆是哪一種（碩／博士｜大學｜專科）：整列只有最左邊那格有字，
                # 列首與每一格頭上都是表頭。這種表不是由上往下一筆一筆填，是照列首放
                # ——資料裡的大學放在「大學」那一列，碩士那一列空著
                keyed = (len(real) >= 3 and len(labels) == 1 and labels[0] is real[0]
                         and len(_squash(real[0][3])) <= 12 and not re.search(r"[：:]", real[0][3])
                         and all(x[4][1] in header_rows for x in real))
                # 空著的、只印格式的（「自　年　月」）、整格勾選框的（「□畢 □肄」）
                # 都是拿來填的。左邊整排都沒有欄位名稱（或只有列首）時，標出它在表頭
                # 底下第幾列——模型靠這個分辨該填第幾筆，沒有它會整批錯位一列
                data = [x for i, x in enumerate(in_row)
                        if x[1] and not is_label(x[3]) and x[4][1] >= 0
                        and (keyed or not any(y[1] and is_label(y[3]) for y in in_row[:i]))]
                # 表頭分兩層時（「區部」底下「日間｜夜間」），第幾列從最下面那層算起
                bottom = max((x[4][1] for x in data), default=-1)
                top = min((x[4][1] for x in data), default=-1)
                seq = 0
                for i, (c, addr, paras, text, head, c0, c1) in enumerate(in_row):
                    if addr is None:
                        continue
                    if re.fullmatch(r"\d{1,2}", text):
                        seq = int(text)         # 序號欄：印著 1、2、3 的那一格
                    # 左邊最近一格印的字才是這一格的欄位名稱。取「整列第一格」會
                    # 把「性別｜　｜血型｜　」的血型欄也標成性別——履歷表一列排
                    # 兩三組「標籤＋值」是常態，錯的標籤比沒有標籤還糟
                    left = next((x[3] for x in reversed(in_row[:i]) if x[3]), "")
                    alone = any(x[1] == addr for x in data)
                    if not alone:       # 標籤＋值排在同一列的格子，頭上的字多半是別的欄位
                        head = _head(heads, edges, c, None, c0, c1 - c0)
                    depth = r - bottom if alone else 0
                    # 說明要緊貼著欄名列才算數，隔了幾列的就不是在講這一區
                    caption = cap_text if depth and cap_row == r - depth - 1 else ""
                    parent = (_head([h for h in heads if h[0] < head[1] and h[0] in header_rows],
                                    edges, -1, -1, head[2], head[3] - head[2])[0]
                              if alone and head[1] in header_rows else "")
                    out.append(Cell(addr, paras, row_head=left, depth=depth,
                                    col_head="" if is_label(text) else head[0],
                                    printed_in_row=len(labels), caption=caption,
                                    value_after=any(x[1] and not x[3] for x in in_row[i + 1:]),
                                    head_row=head[1], col_parent=parent,
                                    origin=top if alone else -1,
                                    key=real[0][3] if keyed and alone else "",
                                    seq=seq if alone else 0,
                                    # 「聯絡電話：(　)↵手機：」夾在「緊急聯絡人｜關係」那一列，
                                    # 左邊最近的欄名只看得到「關係」，看不出是誰的電話
                                    row_first=(real[0][3]
                                               if len(real) >= 3 and len(labels) == len(real)
                                               and real[0][1] != addr and real[0][3] != left
                                               and any(w in _squash(real[0][3])
                                                       for ws in OTHER_PEOPLE.values() for w in ws)
                                               else "")))
                if len(real) >= 3 and len(labels) == len(real):
                    header_rows.add(r)
                # 列首那一列不是表頭：「專科」不是下一列「高中(職)」的欄名
                if labels and not keyed:
                    heads.append((r, len(tcs), [(x[0], x[5], x[6], x[3]) for x in labels]))
            table_no += 1
        else:
            if APPLICANT_AGAIN_RE.search(block.text):
                company = False         # 這一段本身是說明，不是填寫位置
            elif COMPANY_ONLY_RE.search(block.text):
                company = True
            elif not company:
                out.append(Cell(f"p{para_no}", [block]))
            para_no += 1
    return out


def _line_slots(line: str) -> List[Tuple[str, int, int, str]]:
    """一行字裡有哪些可以寫字的位置：(kind, start, end, 選項字或單位)。

    有勾選框的行，框與框之間的空白多半是排版（「□隨時  □   週」），填進去只會把
    版面弄壞。算數的只有兩種：最後一個框之後的（「□   月   日」要寫月份與日期），
    以及夾在同一個選項的字中間的（「□ NT$＿＿＿/月」——選項本身就是一個填空）。
    """
    if not line.strip():
        return []
    squashed = _squash(line)
    if (any(b in squashed for b in SIGN_WORDS + ("年制",))
            or COMPANY_WORDS_RE.search(squashed)):
        return []    # 公司自己填的欄位、親筆簽名、年制
    out: List[Tuple[str, int, int, str]] = []
    boxes = [m.start() for m in re.finditer(f"[{CHECKBOX_CHARS}]", line)]
    for i, pos in enumerate(boxes):
        stop = boxes[i + 1] if i + 1 < len(boxes) else len(line)
        option = re.split(r"[\s　,，、/／_＿]+", line[pos + 1:stop].strip())[0].strip(" :：()（）")
        out.append(("box", pos, pos + 1, option))
    # 「學  歷」「姓    名」中間的空白是把字撐開的排版，不是留給人寫字的。
    # 拿掉空白後只剩三兩個中文字、又沒有冒號、數字或單位的，整段就是個欄位名稱
    # （「　　　,　　　」只剩一個逗號，那是「英文名, 姓氏」兩個要填的空，不是欄位名稱）
    padded = (len(_squash(line)) <= 4 and re.search(r"[一-鿿]", line)
              and not re.search(r"[：:\d]", line) and not any(u in line for u in UNITS))
    if not padded:
        # 行首直接印著單位（「年　月－　年　月」「公分」）：數字寫在單位前面，
        # 那裡沒有留白可以認，位置就是行首
        lead = len(line) - len(line.lstrip(" 　"))
        if lead < 2 and _unit_at(line, lead):
            out.append(("gap", 0, lead, _unit_at(line, lead)))
        for m in GAP_RE.finditer(line):
            if boxes and m.start() < boxes[-1]:
                owner = [b for b in boxes if b < m.start()]
                stop = min(b for b in boxes if b > m.start())
                if (not owner or not line[owner[-1] + 1:m.start()].strip()
                        or not line[m.end():stop].strip()):
                    continue    # 框與框之間、又不是夾在選項的字中間的空白是排版
            unit = _unit_at(line, m.end() + len(line[m.end():]) - len(line[m.end():].lstrip(" 　")))
            filler = m.group()
            if len(filler) < 3 and not unit and filler[0] not in "_＿.．":
                continue    # 兩格寬的空白要後面接著單位才算（「cm  體重」中間那兩格不算）
            if (not line[:m.start()].strip() and not unit
                    and not re.match(r"[,，、/／]", line[m.end():])):
                continue    # 行首的空白後面接欄位名稱就只是縮排（「　　　初試日期：」）
            # 括號裡的留白是區碼（「聯絡電話：(　　)」），號碼其餘的部分寫在括號後面
            if (m.start() and line[m.start() - 1] in "(（"
                    and line[m.end():m.end() + 1] in (")", "）")
                    and line[m.end():m.end() + 1]):
                out.append(("gap", m.start(), m.end() + 1, "()"))
                continue
            out.append(("gap", m.start(), m.end(), unit))
    # 冒號結尾＝值接在後面；「語言:1.」這種編號結尾也是（1. 是清單序號，值寫在後面）。
    # 要看未經 strip 的原字：「中文：      」的留白已經被上面認成 gap 了，
    # 再補一個 append 就會有兩個位置搶同一個地方
    if line.endswith(("：", ":")) or re.search(r"\d[.．、)）]$", line):
        out.append(("append", len(line), len(line), ""))
    elif not out and re.search(r"[：:][ 　]*同[^\s：:]{1,5}[ 　]*$", line):
        # 「通訊地址：同上」：冒號後面印好的「同上」只是提示，地址照樣寫在後面
        # （跟 AT-1 通訊地址欄印著「同戶籍地址」一樣）
        out.append(("append", len(line.rstrip()), len(line.rstrip()), ""))
    elif not out and line.rstrip().endswith(("？", "?")):
        out.append(("line", len(line), len(line), ""))   # 整格是一個問句，答案寫下一行
    return _mark_job_specific(line, out)


def _per_job_key(text: str) -> str:
    """這段字是「這次應徵」的哪一個欄位（應徵職務→job.title），不是就回空字串。"""
    squashed = _squash(text)
    return next((key for word, key in PER_JOB_LABELS.items() if word in squashed), "")


def _mark_job_specific(line: str, slots: List[Tuple[str, int, int, str]]
                       ) -> List[Tuple[str, int, int, str, str]]:
    """標出「應徵職務：＿＿　希望待遇：＿＿」裡哪一段是應徵職務、哪一段是希望待遇：
    只有應徵職務那一段算「這次應徵」，希望待遇照常填。

    一個位置歸哪個欄名，看它前面到上一個位置之間印的字；勾選框、以及框後面的填空
    （「□其他＿＿」）沿用同一題第一個框的歸屬，直到出現新的欄名（冒號）。
    回傳照原本的順序：位置的代碼跟順序有關，順序一動，學過的格式就對不上了。
    """
    if not any(b in _squash(line) for b in JOB_SPECIFIC):
        return [(k, s, e, o, "") for k, s, e, o in slots]
    marked, prev_end, prev_kind, key = {}, 0, "", ""
    for slot in sorted(slots, key=lambda s: s[1]):
        kind, start, end, _option = slot
        region = line[prev_end:start]
        if not (prev_kind == "box" and not re.search(r"[：:]", region)):
            key = _per_job_key(region)
        marked[slot] = key
        prev_end, prev_kind = end, kind
    return [(k, s, e, o, marked[(k, s, e, o)]) for k, s, e, o in slots]


def _window(text: str, at: int, span: int = 34) -> str:
    """▁ 前後各取一段。長的聲明文字只截開頭的話，▁ 附近印了什麼根本看不到。"""
    lo, hi = max(0, at - span), at + span + 1
    return ("…" if lo else "") + text[lo:hi].strip() + ("…" if hi < len(text) else "")


def _fill_after(cell: Cell, text: str) -> bool:
    """這一段印的字後面就是要寫字的地方嗎？

    兩種：整段被括號包起來的寫法說明（戶籍地址欄印著「(請註明里、鄰)」，人就寫在那格），
    以及左邊有欄位名稱、右邊又沒有空格子的短提示（通訊地址欄印著「同戶籍地址」）。
    欄名列不算——一列印著五六個欄位名稱那是表頭，不是提示。
    """
    body = text.strip()
    if re.fullmatch(r"[（(][^（()）]*[)）]", body):
        # 整格就只有這句說明才算。「專 科↵(二.三.五專)」的括號是欄位名稱的補充，不是寫字的地方
        return _squash("".join(p.text for p in cell.paras)) == _squash(body)
    return (bool(cell.row_head) and not cell.value_after and cell.printed_in_row <= 2
            and 0 < len(_squash(body)) <= 8 and not re.search(r"[：:？?]", body))


def slots_of(cell: Cell) -> List[Slot]:
    """一個格子裡所有可以寫字的位置。"""
    out: List[Slot] = []
    full, offset = "\n".join(p.text for p in cell.paras), 0
    # 聲明條文那種長段落之間的空行是排版，不是填寫位置
    legal = any(len(p.text) > 60 and p.text.rstrip().endswith("。") for p in cell.paras)
    first = cell.paras[0].text if cell.paras else ""
    for pi, para in enumerate(cell.paras):
        text = para.text
        above = cell.paras[pi - 1].text.rstrip() if pi else ""
        if PLACEHOLDER_RE.match(text):
            # 「中文姓名：↵（空行）」的空行是上一行那個欄位的地方：值接在冒號後面，
            # 冒號那一行被擋掉的（應徵職務：）空行也跟著不填
            found = ([] if legal or above.endswith(("：", ":"))
                     else [("blank", 0, len(text), "", "")])
        else:
            found, base = [], 0
            for line in text.split("\n"):
                found += [(k, s + base, e + base, opt, job)
                          for k, s, e, opt, job in _line_slots(line)]
                base += len(line) + 1
            if not found and _fill_after(cell, text):
                found = [("append", len(text), len(text), "", "")]
            # 「聯絡電話：(　　)↵手機」：第二段只印一個欄位名稱，號碼寫在它後面
            if (not found and pi and re.search(r"[：:]", first)
                    and _squash(text) in set(LABEL_ALIASES) | {f.label for f in BY_KEY.values()}):
                found = [("append", len(text.rstrip()), len(text.rstrip()), "", "")]
        # 「這次應徵」的位置另外編號（#0.j1）：照原本的順序插進去會把後面的編號往後推，
        # 學過的格式就整份對不上了
        n = j = 0
        for kind, start, end, option, job_field in found:
            if job_field:
                j += 1
            else:
                n += 1
            shown_end = end - 1 if option == "()" else end
            if len(cell.paras) > 1 and (kind == "blank" or not re.search(r"[一-鿿A-Za-z]",
                                                                        text[:start])):
                # 格子裡的空白段本身沒有字，行首的空格（「▁年　月　日」「▁公分」）前面也
                # 沒有字，要把整格攤開才看得出它屬於哪個欄位
                shown = full[:offset + start] + "▁" + full[offset + shown_end:]
            else:
                shown = text[:start] + "▁" + text[shown_end:]
            # 排版用的留白先縮成一格：宣告事項那六題的題目與勾選框之間空了五十格，
            # 不縮的話周圍只看得到「▁ 是,請說明:」，題目在問什麼完全看不見
            shown = re.sub(r"[ 　]{2,}", " ", shown.replace("\n", "↵"))
            out.append(Slot(f"{cell.addr}#{pi}.{'j' if job_field else ''}{j if job_field else n}",
                            kind, cell.addr, pi, start, end,
                            filler=text[start:end] if kind == "gap" else "",
                            option=option, preview=_window(shown, shown.index("▁")),
                            job_field=job_field, extra=bool(job_field), cell=cell))
        offset += len(text) + 1
    # 格子裡印了「優點：」「缺點：」這種更明確的位置時，問句後面就不算一個位置了——
    # 不然模型會把答案寫在問句後面，印好的欄位反而空著
    if any(s.kind == "append" for s in out):
        out = [s for s in out if s.kind != "line"]
    # 問句的下一段一開頭就是勾選框（「如何得知該職缺資訊？↵□人力銀行 □親友推薦」）：
    # 答案是勾，不是在問句底下另寫一行字
    out = [s for s in out if not (
        s.kind == "line" and s.para + 1 < len(cell.paras)
        and re.match(f"[ 　]*[{CHECKBOX_CHARS}]", cell.paras[s.para + 1].text))]
    # 「出生日期：↵年　月　日」「身高：↵公分」：下一段印著單位，值照單位寫在那裡，
    # 冒號後面不再是位置——不然整串日期寫在冒號後面，底下的年月日照樣空著
    united = {s.para for s in out if s.kind == "gap" and s.option in UNITS}
    out = [s for s in out if not (s.kind == "append" and s.para + 1 in united
                                  and s.end == len(cell.paras[s.para].text))]
    # 整格空白、欄名印在隔壁那格（「應徵職務｜＿＿」）：這一格就是這次應徵的職務。
    # 只認整格空白的：row_head 是「左邊最近印著字的那格」，自己有印字的格子
    # （「應徵職務：＿｜錄取後可報到日 □隨時」）左邊那個欄名是別人的，認了就會搶錯格
    if all(not p.text.strip() for p in cell.paras):
        label = _per_job_key(cell.row_head) or _per_job_key(cell.col_head)
        for slot in out:
            if label and not slot.job_field:
                slot.job_field = label
    return out


# ---------------------------------------------------------------------------
# 二、個人資料攤平成「項目代碼 → 值」
# ---------------------------------------------------------------------------

def flatten(profile: Any, prefix: str = "") -> Dict[str, str]:
    out: Dict[str, str] = {}
    if isinstance(profile, dict):
        for key, value in profile.items():
            out.update(flatten(value, f"{prefix}.{key}" if prefix else key))
    elif isinstance(profile, list):
        for i, value in enumerate(profile):
            out.update(flatten(value, f"{prefix}[{i}]"))
    elif str(profile).strip():
        out[prefix] = str(profile).strip()
    return out


def _spec(path: str):
    """education[0].school → schema 裡 education[].school 的定義。"""
    return BY_KEY.get(re.sub(r"\[\d+\]", "[]", path))


def _label(path: str) -> str:
    spec = _spec(path)
    return f"（{spec.label}）" if spec else ""


def _kind(path: str) -> str:
    spec = _spec(path)
    return spec.kind if spec else ""


# 整個值就是一個日期：2016/9、2013年09月、2023 年 7 月、1998年03月25日、民國85年3月 都算
_DATE_ONLY_RE = re.compile(
    r"^\s*(民國)?\s*(\d{2,4})\s*[年/.\-]\s*(\d{1,2})\s*[月/.\-]?\s*(?:(\d{1,2})\s*日?)?\s*$")


def _canon_date(value: str) -> str:
    """日期一律寫成「2016年9月」。

    使用者存進來的寫法從來沒一致過——同一份 app.db 裡就有 2016/9、2013年09月、
    2023 年 7 月三種。原樣照抄的話，學歷表上下兩列會長得不一樣。

    印著「＿年＿月」的格子不受影響：那種格子是照印好的單位一格一格填，年月由
    表格提供，這裡補上的年月反而會被格子的單位重複。認不出是日期就原樣保留，
    不替使用者猜——「民國87年」「2020」這種留給它原本的樣子。
    明寫民國的（「民國85年3月」，多半是匯入的）換成西元：內部一律西元，表格要民國時
    寫的那一端再換回去（_roc）。沒寫民國的兩三位數年份不猜——「85/3」也可能是西元 1985。
    """
    m = _DATE_ONLY_RE.match(value or "")
    if not m:
        return value
    roc, year, month, day = m.groups()
    year = int(year) + (1911 if roc and int(year) < 1911 else 0)
    return f"{year}年{int(month)}月" + (f"{int(day)}日" if day else "")


def fields_of(profile: Dict[str, Any]) -> Dict[str, str]:
    """產品的個人資料（app.db 那份 JSON）攤平成「項目代碼 → 值」。

    日期一律換成同一種寫法（見 _canon_date）：存的是使用者自己打的，填出去的
    要是一致的，換寫法是填寫這一端的事。

    schema 標成合成的欄位（就學期間、任職期間）不採用存著的值，一律由起訖重算——
    app.db 裡就有一筆存成「自 2023 年 2 月 至 2021 年 10 月」，起訖顛倒。
    """
    out = {k: (_canon_date(v) if _kind(k) == "date" else v)
           for k, v in flatten(profile).items()
           if not getattr(_spec(k), "derived", False)}
    for root in ("education", "experience"):
        for i, row in enumerate(profile.get(root) or []):
            if isinstance(row, dict) and row.get("start") and row.get("end"):
                out[f"{root}[{i}].period"] = (f"{_canon_date(row['start'])}"
                                              f"~{_canon_date(row['end'])}")
                if root == "experience" and _tenure(row["start"], row["end"]):
                    out[f"{root}[{i}].tenure"] = _tenure(row["start"], row["end"])
    basic = profile.get("basic") or {}
    # 生日曆制＝民國：存的「85年04月15日」是民國年，換成西元（見 _canon_date）
    birthday = str(basic.get("birthday") or "").strip()
    if basic.get("birthday_era") == "民國" and birthday and not birthday.startswith("民國"):
        out["basic.birthday"] = _canon_date("民國" + birthday)
    surname = _surname_en(str(basic.get("name_passport") or ""), str(basic.get("name_zh") or ""))
    if surname:
        out["basic.surname_en"] = surname
    # 年齡由生日算到今天：存著的會過期，去年填的今年還是去年的歲數
    age = _age(out.get("basic.birthday", ""))
    if age:
        out["basic.age"] = age
    # 服役期間與總年資同理：起訖或工作經歷改了就跟著變
    if basic.get("military_start") and basic.get("military_end"):
        out["basic.military_period"] = (f"{_canon_date(str(basic['military_start']))}"
                                        f"~{_canon_date(str(basic['military_end']))}")
    # 總年資：有寫起訖卻算不出來的那一筆，不能當成 0 個月偷偷少算——那會寫出一個
    # 比實際短的年資，比空著更糟（同 _tenure 的原則）。整筆沒寫起訖的不算數：
    # 那是使用者沒有主張期間，不是算不出來
    dated = [row for row in (profile.get("experience") or [])
             if isinstance(row, dict)
             and str(row.get("start") or "").strip() and str(row.get("end") or "").strip()]
    spans = [_months(str(row["start"]), str(row["end"])) for row in dated]
    if spans and all(spans):
        out["basic.total_tenure"] = _span(sum(spans))
    return out


def _age(birthday: str) -> str:
    """幾歲：生日還沒到就少一歲。認不出生日就不猜。"""
    m = DATE_RE.search(birthday or "")
    if not m:
        return ""
    today = _today()
    year, month = int(m.group(1)), int(m.group(2))
    day = int(m.group(3)) if m.lastindex and m.lastindex >= 3 and m.group(3) else 1
    age = today.year - year - ((today.month, today.day) < (month, day))
    return str(age) if 0 < age < 120 else ""


def _today() -> date:
    return date.today()


def _months(start: str, end: str) -> int:
    """這段期間有幾個月，頭尾都算。認不出來或顛倒就回 0。

    認不出來的兩條路以前回的是空字串，跟宣告的 int 對不上：呼叫端一個做
    `divmod(months, 12)`、一個做 `sum(...)`，使用者只要在工作經歷打「2018年」
    這種只有年份的寫法（`DATE_RE` 要年＋月，而 `_canon_date` 刻意不替使用者猜
    年份-only 的日期），`fields_of` 就整個拋 TypeError——而它在分析與匯出都會跑，
    等於每一次都當掉。
    """
    a = DATE_RE.search(start or "")
    if PRESENT_RE.match(end or ""):
        today = _today()
        end_year, end_month = today.year, today.month
    else:
        b = DATE_RE.search(end or "")
        if not b:
            return 0
        end_year, end_month = int(b.group(1)), int(b.group(2))
    if not a:
        return 0
    months = (end_year * 12 + end_month) - (int(a.group(1)) * 12 + int(a.group(2))) + 1
    return months if months > 0 else 0


def _span(months: int) -> str:
    years, rest = divmod(months, 12)
    return (f"{years}年" if years else "") + (f"{rest}個月" if rest else "")


def _tenure(start: str, end: str) -> str:
    """年資：頭尾兩個月都算，跟 104、LinkedIn 的算法一樣（2023年7月～2026年4月＝2年10個月）。
    還在職（訖是「至今」）就算到這個月。
    認不出日期、或起訖顛倒就不算——寧可空著，不寫一個錯的年資。"""
    return _span(_months(start, end))


# 國語羅馬拼音的音節（威妥瑪、漢語、通用拼音混著收）。護照全名存成「KUOWEITE」這種
# 沒有分隔的寫法時，要靠它切出「KUO WEI TE」才知道姓氏是哪一段
_SYLLABLE_RE = re.compile(
    r"(?:CH|SH|ZH|TS|TZ|HS|SZ|SS|[BPMFDTNLGKHJQXZCSRWY])?"
    r"(?:IUNG|IANG|IONG|UANG|UENG|ANG|ENG|ING|ONG|UNG|IAN|IAO|IEN|IEH|UAI|UAN|UEI|UEN|UEH"
    r"|AI|AO|AN|EI|EN|ER|EH|IA|IE|IH|IN|IO|IU|OU|UA|UE|UI|UN|UO|U|A|E|I|O)")


def _surname_en(passport: str, name_zh: str) -> str:
    """護照全名裡的姓氏拼音。

    有分隔的（「KUO, WEI-TE」「KUO WEI TE」）取第一段；只有兩段而其中一段帶連字號
    （「WEI-TE KUO」），帶連字號的是名字，另一段才是姓。沒有分隔的（「KUOWEITE」）
    照中文姓名的字數切音節——三個字就要剛好切成三個音節，所有切法的第一個音節都
    相同才採用，切得出兩種姓氏就不猜。四個字的名字當成複姓（歐陽、司馬）。
    """
    text = passport.strip().upper()
    parts = [p for p in re.split(r"[\s,，]+", text) if p]
    if len(parts) == 2 and sum("-" in p for p in parts) == 1:
        return next(p for p in parts if "-" not in p)
    if len(parts) > 1:
        return parts[0]
    count = len(re.findall(r"[一-鿿]", name_zh))
    if not parts or not re.fullmatch(r"[A-Z]+", parts[0]) or not 2 <= count <= 4:
        return ""
    word, heads = parts[0], 2 if count == 4 else 1

    def split(rest: str, n: int) -> List[List[str]]:
        if not rest or not n:
            return [[]] if not rest and not n else []
        out = []
        for size in range(1, min(6, len(rest)) + 1):
            if _SYLLABLE_RE.fullmatch(rest[:size]):
                out += [[rest[:size]] + tail for tail in split(rest[size:], n - 1)]
        return out

    surnames = {"".join(s[:heads]) for s in split(word, count)}
    return surnames.pop() if len(surnames) == 1 else ""


# ---------------------------------------------------------------------------
# 三、版面示意圖
# ---------------------------------------------------------------------------

def _wrap(draw, text: str, font, width: float) -> List[str]:
    """照像素寬度斷行。中文沒有空白可以斷，只能逐字量。"""
    out: List[str] = []
    for raw in text.split("\n"):
        line = ""
        for ch in raw:
            if line and draw.textlength(line + ch, font=font) > width:
                out.append(line)
                line = ch
            else:
                line += ch
        out.append(line)
    return out or [""]


def _layout(doc, draw, font) -> tuple:
    """算出每一格畫在哪裡。回傳 ([(方框, 地址, 文字行)], [(y, 地址)], 總高度)。

    欄寬照 docx 自己的 tblGrid（表格真正的欄位寬度），橫向合併用 gridSpan 併欄，
    所以畫出來的相對位置與寬窄跟原檔一致。

    第二份清單是空白段落：它們畫不出東西，但仍然是可寫位置，得知道排在哪一頁——
    不記的話，凡是分到這種位置的批次都查不到頁碼，只好整份圖附上。
    """
    items, ghosts, y = [], [], 10
    table_no = para_no = 0
    for block in iter_block_items(doc):
        if isinstance(block, Table):
            cols = [int(g.w or 0) for g in block._tbl.tblGrid.gridCol_lst] or [1]
            edges, acc = [0.0], 0
            for w in cols:
                acc += w
                edges.append(acc)
            xs = [10 + (PAGE_W - 20) * e / (acc or 1) for e in edges]
            for r, row in enumerate(block.rows):
                grid_col, in_row = 0, []
                for c, tc in enumerate(row._tr.tc_lst):
                    x0 = xs[min(grid_col, len(xs) - 1)]
                    grid_col += tc.grid_span or 1
                    x1 = xs[min(grid_col, len(xs) - 1)]
                    text = "\n".join(p.text for p in _Cell(tc, block).paragraphs).strip()
                    in_row.append((x0, x1, f"t{table_no}.r{r}.c{c}",
                                   _wrap(draw, text, font, max(x1 - x0 - 10, 20))))
                height = max(len(lines) for *_, lines in in_row) * LINE_H + ADDR_H + 6
                items += [((x0, y, x1, y + height), addr, lines)
                          for x0, x1, addr, lines in in_row]
                y += height
            table_no += 1
            y += 8
        else:
            if block.text.strip():
                lines = _wrap(draw, block.text.strip(), font, PAGE_W - 30)
                height = len(lines) * LINE_H + ADDR_H + 6
                items.append(((10, y, PAGE_W - 10, y + height), f"p{para_no}", lines))
                y += height
            else:
                ghosts.append((y, f"p{para_no}"))
            para_no += 1
    return items, ghosts, y + 10


@dataclass
class Pages:
    """版面示意圖，外加「每一格畫在第幾張上」。

    圖片是提示快取的分界線：實測同一份提示重送只重算 4 個 token，但只要換掉結尾
    那一段，4431 個 token 裡就有 4026 個要重算——能重用的只有第一張圖前面那段文字。
    加 `--cache-reuse` 也一樣。所以每一次呼叫附幾張圖，就是實實在在的時間。

    配對、一列一筆、補漏三輪只附這一批位置所在的那幾張（`for_addrs`）——那幾輪要的
    是「找得到位置」。勾選題整份附上：那一輪要判斷「這題在問哪一項資料」，少看一張
    會改判（連跑兩輪都一樣），見 `ask_boxes` 的註解。
    """

    urls: List[str]                        # 已經編成 data URI，省掉每一次呼叫重編
    page_of: Dict[str, Tuple[int, int]]    # 地址 -> (起頁, 迄頁)，跨頁的格子佔兩張

    def __len__(self) -> int:
        return len(self.urls)

    def all(self) -> List[int]:
        return list(range(len(self.urls)))

    def for_addrs(self, addrs) -> List[int]:
        """這一批位置所在的那幾張，回傳張數編號。

        地址查不到就整份附上：寧可慢一次，也不要讓模型對著沒有那一格的圖硬猜
        （`_layout` 畫的地址是 `cells()` 的超集，正常情況下不會查不到）。
        """
        want: set = set()
        for addr in addrs:
            span = self.page_of.get(addr)
            if span is None:
                return self.all()
            want.update(range(span[0], span[1] + 1))
        return sorted(want) or self.all()

    def parts(self, wanted: List[int]) -> List[Dict[str, Any]]:
        """這幾張的訊息內容。"""
        return [{"type": "image_url", "image_url": {"url": self.urls[i]}} for i in wanted]


def render_pages(doc) -> Pages:
    """把版面畫成示意圖，一張圖一頁。

    不做像素級還原（那需要排版引擎，而 LibreOffice、Word 都不能用——產品是可攜
    離線的單機工具）。模型需要的本來也不是字型，而是「哪一格在哪、多寬、跟誰同一列」，
    這些從 docx 的表格骨架就畫得出來。

    每一格的地址直接印在格子左上角：模型看得到位置，才對得回它要回答的地址。
    順便記下每一格落在第幾張，讓每一批只附自己看得到的那幾張。
    """
    font = ImageFont.truetype(FONT_PATH, FONT_SIZE)
    small = ImageFont.truetype(FONT_PATH, ADDR_SIZE)
    items, ghosts, height = _layout(doc, ImageDraw.Draw(Image.new("RGB", (1, 1))), font)

    canvas = Image.new("RGB", (PAGE_W, height), "white")
    draw = ImageDraw.Draw(canvas)
    for (x0, y0, x1, y1), addr, lines in items:
        draw.rectangle((x0, y0, x1, y1), outline=(150, 150, 150))
        draw.text((x0 + 4, y0 + 1), addr, font=small, fill=(190, 60, 60))
        for i, line in enumerate(lines):
            draw.text((x0 + 5, y0 + ADDR_H + i * LINE_H), line, font=font, fill=(0, 0, 0))

    scale = 1.0
    if height > MAX_PAGES * PAGE_H:      # 太長的表格整張縮小，不然圖多到塞爆上下文
        scale = MAX_PAGES * PAGE_H / height
        canvas = canvas.resize((PAGE_W, MAX_PAGES * PAGE_H))
        height = MAX_PAGES * PAGE_H

    urls = []
    for top in range(0, height, PAGE_H):
        buf = io.BytesIO()
        canvas.crop((0, top, canvas.width, min(top + PAGE_H, height))).save(buf, "PNG")
        urls.append("data:image/png;base64,"
                    + base64.b64encode(buf.getvalue()).decode("ascii"))

    # 縮過的話座標要跟著縮。格子底邊剛好壓在切線上算前一張（1500 是第 0 張的最後一列）
    last = len(urls) - 1
    page_of = {}
    for (_x0, y0, _x1, y1), addr, _lines in items:
        lo = min(int(y0 * scale) // PAGE_H, last)
        hi = min(max(int(y1 * scale - 1) // PAGE_H, lo), last)
        page_of[addr] = (lo, hi)
    for y, addr in ghosts:
        page = min(int(y * scale) // PAGE_H, last)
        page_of[addr] = (page, page)
    return Pages(urls, page_of)


# ---------------------------------------------------------------------------
# 四、問模型：這個位置放哪一項資料
# ---------------------------------------------------------------------------

def _describe(slot: Slot) -> str:
    cell, parts = slot.cell, []
    if cell.row_head:
        parts.append(f"左邊欄名:{_squash(cell.row_head)[:14]}")
    if cell.col_head:
        depth = f"（往下第{cell.depth}列）" if cell.depth else ""
        parts.append(f"上面欄名:{_squash(cell.col_head)[:14]}{depth}")
    if cell.row_first:
        parts.append(f"這一列開頭:{_squash(cell.row_first)[:10]}")
    kind = (f"（勾選框，選項印著「{slot.option}」）" if slot.kind == "box"
            else "（括號裡寫區碼，其餘號碼接在括號後面，填市話）" if slot.option == "()" else "")
    return f"{slot.preview}{kind}" + "".join(f"｜{x}" for x in parts)


def ask(batch: List[Slot], fields: Dict[str, str], pages: Pages,
        used: Dict[str, str], chat: llm.Chat) -> Dict[str, str]:
    """問模型這一批位置各自要填哪一項資料。回傳 {位置編號: 項目代碼}。"""
    # 位置編號當成 JSON 的鍵：文法會自己把編號一個一個吐出來，模型只填「哪一項資料」。
    # 之前讓模型自己配 {slot, field} 一對一對地寫，它會整批位移一格
    # （優缺點那格把「優點」寫進問句行、「缺點」寫進優點欄）。鍵由文法產生就錯不了。
    # 鍵用 s1、s2 這種短名，不用完整地址——地址一個 12 個字，16 個位置光是把鍵
    # 逐字吐出來就佔掉大半輸出，模型剩不了多少心力在「挑哪一項」上
    ids = [f"s{i}" for i in range(1, len(batch) + 1)]
    choices = {"type": "string", "enum": list(fields) + [NONE]}
    schema = {
        "type": "object",
        "properties": {sid: choices for sid in ids},
        "required": ids,
        "additionalProperties": False,
    }
    # 這一輪的幾批接在同一串對話上問：個人資料與示意圖只有第一次用到時附上，後面幾批
    # 只接新的問題，llama-server 的提示快取才吃得到。也不重貼整份位置清單——整張表格
    # 的長相示意圖已經畫給模型看了，再貼一份 180 行的清單只是讓提示多五千個 token。
    # 示意圖只附這一批位置所在的那幾張：分批本來就照文件順序切，一批十六個位置
    # 幾乎都落在同一張上（見 Pages）
    #
    # attached 收齊這一次附了什麼，交給 chat.ask 在問成之後才記帳：這裡到 ask 之間
    # 要是出了事，這些就不算附過，下一批會重新附上
    chat.start_over_if_long()
    attached: List[str] = []
    user: List[Dict[str, Any]] = []
    if not chat.seen("fields"):
        attached.append("fields")
        user.append({"type": "text", "text": (
            "個人資料（項目代碼：值）：\n"
            + "\n".join(f"- {k}{_label(k)}：{v[:60]}" for k, v in fields.items()))})
    fresh = [i for i in pages.for_addrs({s.addr for s in batch}) if not chat.seen(f"page{i}")]
    attached += [f"page{i}" for i in fresh]
    user += pages.parts(fresh)
    # 已經用掉的資料排在最後（放前面會把提示快取的共同前綴打斷）。沒有這一段的話，
    # 模型看不到別批的決定：五個問答題配七項回答，每一批都從頭挑一次就會互相搶
    done = ("\n\n已經填在別的位置的資料項（不要再挑，除非這一格真的也要填同一項）：\n"
            + "、".join(sorted(set(used.values()))) if used else "")
    user.append({"type": "text", "text": "這一批要判斷的位置：\n" + "\n".join(
        f"  {sid}｜{_describe(s)}" for sid, s in zip(ids, batch)) + done})

    data = chat.ask(user, schema, label=f"配對:{len(ids)}處", attached=attached)
    back = dict(zip(ids, batch))
    return {back[sid].id: key for sid, key in data.items()
            if sid in back and key in fields}


# ---------------------------------------------------------------------------
# 五、把值寫進去。原本印的字一個都不動
# ---------------------------------------------------------------------------

def _squash(text: str) -> str:
    return re.sub(r"[\s　]+", "", text or "")


def _ticked(option: str, value: str) -> bool:
    """這個選項該不該勾：選項印的字要跟值（或值裡用頓號分開的其中一項）完全相同。

    刻意不做包含比對——「同意」是「不同意」的子字串，包含比對會兩個都勾起來。
    反過來「Windows、Word、Excel」拆開後每一項都對得上，多選也照樣成立。
    """
    if not option:
        return False
    cands = {_squash(value)} | {_squash(p) for p in OPTION_SPLIT_RE.split(value) if p}
    return _squash(option) in cands


def _marker_after(text: str, end: int) -> str:
    """空格後面緊接著印的單位：「民國　▁　年」的 ▁ 後面是「年」。"""
    return _unit_at(text[end:].lstrip(" 　"), 0)


def _split_by_markers(value: str, markers: List[str]) -> List[str]:
    """把一個值照單位切開：「1998年03月25日」＋［年,月,日］→［1998, 3, 25］。

    每一段取單位前面最後一個數：「2026 年 8 月 24日」只照［月,日］切時，
    月那一段是「2026 年 8」，要的是 8——表格沒印年，年份就不寫。
    數要取完整的：「75,000元」是 75,000、「170.5公分」是 170.5（以前只取最後一組
    數字，分別剩下 0 與 5）。整數的前導零拿掉，人寫「3 月」不寫「03 月」。

    切完必須剛好用完整個值，剩下尾巴就算失敗——「1998年03月25日」只照［年,月］
    切得出 1998、3，但 25 日沒地方去，那就不是這個值該去的地方。
    """
    parts, cursor = [], 0
    for marker in markers:
        idx = value.find(marker, cursor)
        if idx < 0:
            return []
        nums = re.findall(r"\d[\d,]*(?:\.\d+)?", value[cursor:idx])
        num = nums[-1].rstrip(",") if nums else ""
        if num.isdigit():
            num = num.lstrip("0") or "0"
        parts.append(num or value[cursor:idx].strip())
        cursor = idx + len(marker)
    return parts if all(parts) and not value[cursor:].strip() else []


def _width(text: str) -> int:
    """印出來佔幾個半形寬：中文字、全形空白算兩個。"""
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in text)


def _pad(slot: Slot, value: str) -> str:
    """值取代整段留白，前後各留一個空白就好——前提是放得進原本的寬度。

    留白是給人手寫的空間，字打上去就不需要了。原本把剩下的空白全留著（值靠左或置中），
    「□ 隨時　□　8　月　24　日後」這種本來就排滿的一行會變長、被擠到下一行
    （使用者親自修答案卷時指出的）。所以只在「值＋前後各一個空白」不超過原本寬度時
    才留空白，放不下就一個都不留。行首、冒號、左括號與破折號後面不留前面那個（「－2021年」），
    逗號、斜線、右括號前面不留後面那個（「75,000/月」）。
    """
    if not slot.filler.isspace():
        return value
    # 半形全形混著的留白（「身高　    cm」）用半形空白隔：全形的太寬，常常放不下就一個都不隔
    sp = " " if " " in slot.filler else slot.filler[0]
    if _width(slot.filler) < _width(value) + 2 * _width(sp):
        return value
    text = slot.cell.paras[slot.para].text
    before, after = text[:slot.start].split("\n")[-1], text[slot.end:]
    padded = sp + value + sp
    if not before.strip() or re.search(r"[(（「［\[－—–~～\-：:]$", before):
        padded = padded.lstrip(" 　")
    if re.match(r"[,，、。：:/／)）」］\]]", after):
        padded = padded.rstrip(" 　")
    return padded


def _area_code(slot: Slot, value: str) -> str:
    """「(　　)」這種區碼位置：區碼寫進括號、其餘號碼接在括號後面。

    位置的範圍包含右括號，所以右括號由這裡原樣寫回去。手機（09 開頭）沒有區碼，
    不是這個位置該填的——回空字串，那一格就不寫。
    """
    text = (value or "").strip()
    digits = re.sub(r"\D", "", text)
    given = re.match(r"\(?(0\d{1,3})\)?[\s\-)]", text)      # 值自己就分好了：02-24596466
    code = (given.group(1) if given
            else next((c for c in AREA_CODES if digits.startswith(c)), ""))
    if not code or digits.startswith("09") or len(digits) < 9:
        return ""
    # 括號裡不留空白：「(02)24596466」，跟一般寫電話的樣子一樣
    return code + slot.filler[-1] + digits[len(code):]


def _replacement(slot: Slot, value: str) -> str:
    if slot.kind == "box":
        # 沒有模型判斷過的勾（或資料改了、舊的勾不算數）時才走到這裡：字面加同義詞表，
        # 「有」對「□是」、「無」對「□否」——以前只比字面，答案改了就兩個都不勾
        return CHECKED if _synonym_hit(slot.option, value) else ""
    text = slot.cell.paras[slot.para].text if slot.cell else ""
    value = _roc_value(slot, text, value)
    if slot.kind == "line":
        return "\n" + value
    if slot.kind == "gap":
        return _pad(slot, value)
    # 接在沒有冒號的字後面（「手機」「(請註明里、鄰)」）隔一格，不然號碼黏著欄名；
    # 冒號與編號的點後面不隔（「語言:1.英文」）
    before = text[:slot.start]
    if slot.kind == "append" and before and not re.search(r"[\s　：:.．、]$", before):
        return " " + value
    return value


def _roc_wanted(slot: Slot, text: str) -> bool:
    """表格要民國年嗎：這個位置前面緊接著「民國」（「民國＿＿年」「出生日期（民國）：＿」），
    或這一格的欄名、表頭印著民國。「中華民國」是國籍，不算。"""
    if ROC_BEFORE_RE.search(text[:slot.start]):
        return True
    return bool(slot.cell and ROC_WORD_RE.search(f"{slot.cell.row_head}{slot.cell.col_head}"))


def _roc(text: str, slot: Slot, part: str, marker: str) -> str:
    """表格要民國年而資料存西元年時，換成民國年。資料存西元是對的——
    表格有的印民國、有的印西元，換算是填寫這一端的事。"""
    if marker == "年" and part.isdigit() and int(part) > 1911 and _roc_wanted(slot, text):
        return str(int(part) - 1911)
    return part


def _roc_value(slot: Slot, text: str, value: str) -> str:
    """整個日期寫進一格時（「出生年月日(民國)」底下那格），表格要民國就整串換。"""
    return to_roc(value) if DATE_RE.search(value) and _roc_wanted(slot, text) else value


def _date_parts(value: str, markers: List[str]) -> List[str]:
    """把日期填進「＿年＿月＿日」這種格子。格子問的全是日期單位時才動手。

    「2016/9」「2013年09月」「自 2021 年 10 月 至 2023 年 2 月」寫法各不相同，照字面
    切是切不開的（值裡根本沒有「年」）。這裡先把值裡的日期解析出來，再照格子問的
    單位一格一格給；遇到第二個「年」就換下一個日期——「自＿年＿月 至＿年＿月」
    要的是起訖兩個。
    """
    if not markers or any(m not in DATE_UNITS for m in markers):
        return []
    found = list(DATE_RE.finditer(value))
    if not found:
        return []
    dates = [m.groups() for m in found]
    # 「2023年7月~至今」：訖那一組的第一格寫「至今」，同組其他格留白（空字串＝不寫）
    ongoing = re.search(PRESENT_WORDS, value[found[-1].end():], re.IGNORECASE)
    out, i, seen = [], 0, set()
    for marker in markers:
        if marker in seen:
            i, seen = i + 1, set()
        if i >= len(dates):
            if not (ongoing and i == len(dates)):
                return []
            out.append("" if seen else ongoing.group())
            seen.add(marker)
            continue
        part = dict(zip(("年", "月", "日"), dates[i])).get(marker)
        if not part:
            return []
        out.append(part)
        seen.add(marker)
    return out


def _spread(run: List[Slot], value: str, texts: Dict[str, str]) -> Dict[str, str]:
    """一項資料填進一串連續的空格：先當日期解析，不行再照空格後面印的單位切。

    「民國　＿　年　＿　月　＿　日」配生日、「自＿年＿月／至＿年＿月」配任職期間都是
    這個形態，而模型往往只指名第一格。所以不是只看被指名的格子，而是從被指名那一格
    往後看整串空格（跨段落也算，「自…」與「至…」是同一格裡的兩段）；
    切不開就只填第一格——整串值塞進第一格比留白還糟。
    """
    if run[0].option == "()":
        written = _area_code(run[0], value)
        return {run[0].id: written} if written else {}
    markers = [_marker_after(texts[x.id], x.end) for x in run]
    parts = _date_parts(value, markers)
    if not parts:
        for k in range(len(run), 1, -1):
            if all(markers[:k]):
                parts = _split_by_markers(value, markers[:k])
                if parts:
                    break
    if not parts:
        return {run[0].id: _pad(run[0], _roc_value(run[0], texts[run[0].id], value))}
    # 空字串的那幾格不寫（「至今」只寫在訖那一組的第一格）——連空白都不能補，
    # 不然留給人手寫的底線會被兩個空白蓋掉
    out = {x.id: _pad(x, _roc(texts[x.id], x, part, marker))
           for x, part, marker in zip(run, parts, markers) if part}
    # 月、日這種一兩位數的幾格要一致：「2021年10月－2023年 2 月」「■ 8 月24日」一格有空白
    # 一格沒有，看起來像打錯。有一格放不下空白就全部不留；四位數的年份本來就常貼著寫，不算
    short = [x.id for x, part in zip(run, parts) if part and len(part) <= 2]
    if len({out[i] != out[i].strip(" 　") for i in short}) > 1:
        out.update({i: out[i].strip(" 　") for i in short})
    return out


def _bigrams(text: str) -> set:
    """連續兩個字的詞，只留含中文、又不是雜訊的。

    「in」（Windows 跟 Linux 都有）、「年月」（離職年月跟服役期間「自＿年＿月」都有）、
    「工作」（工作內容跟「您以前在工作上有什麼貢獻」都有）這種詞兩邊都出現，
    卻不代表相關，拿來比對只會放行配錯的資料。
    """
    t = _squash(text)
    return {b for b in (t[i:i + 2] for i in range(len(t) - 1))
            if re.search(r"[一-鿿]", b) and b not in GENERIC_WORDS
            and not re.fullmatch(r"[年月日時分秒]{2}", b)}


def _around(slot: Slot) -> str:
    """位置周圍印的字：欄名，加上「管著這個位置的那一段」——從前一個冒號或換行到 ▁，
    再帶一小段後面的字。

    整格的字全算進來的話，「役別：▁　(如免役說明原因)：▁」這種一格兩欄的位置，
    兩個空格比對出來的結果一模一樣，配對對調了也看不出來。
    「本公司」「本人」拿掉：那是指發表格的公司與填表人自己，不是資料裡的哪一家公司、
    哪一個人。留著的話，推薦人的公司會因為「公司」兩個字被填進「您對本公司的了解」。
    """
    # 一格分好幾段時，標籤常只印在第一段（「負債狀況：」在第一段，勾選框在第二段）
    # 這一段在位置前面沒印任何字（「▁　　, ▁」「▁公分」）時，第一段就是它的欄位名稱
    # 第一段是「聯絡電話：(　　)」時，第二段的「手機：▁」也是聯絡電話底下的一項——取到冒號為止
    label = ""
    if slot.para > 0:
        first = slot.cell.paras[0].text.strip()
        own = slot.cell.paras[slot.para].text[:slot.start]
        if not re.search(r"[一-鿿A-Za-z]", own):
            # 往前找最近一段不是勾選框開頭的：「是否有配偶…任職？↵ □否 ↵□是，請說明」
            # 的兩個框都屬於第二段那一題，不是第一段那一題
            label = next((p.text.strip() for p in reversed(slot.cell.paras[:slot.para])
                          if p.text.strip()
                          and not re.match(f"[ 　]*[{CHECKBOX_CHARS}]", p.text)), first)
        elif first.endswith(("：", ":")):
            label = first
        elif re.search(r"[：:]", first):
            label = re.match(r".*[：:]", first).group()
    text = slot.preview
    i = text.find("▁")
    if i < 0:
        near = text
    else:
        head = text[:i].rstrip("：: 　")
        cut = max(head.rfind(c) for c in "：:↵。")
        # 冒號後面只印著「同上」「同戶籍地址」這種提示時，欄位名稱在冒號前面（「通訊地址：同上▁」）
        if cut >= 0 and re.fullmatch(r"[ 　]*同[^\s：:]{0,5}[ 　]*", head[cut + 1:]):
            cut = max(head.rfind(c, 0, cut) for c in "：:↵。")
        near = head[cut + 1:] + text[i:i + 8]
    return re.sub(r"本公司|本人", "", slot.cell.row_head + slot.cell.col_head + label + near)


def _aliases(key: str) -> List[str]:
    """產品對照表裡這個欄位的別名：「手機」也是行動電話、「服務單位」也是公司名稱。

    schema.LABEL_ALIASES 是產品累積下來的表格用語，比對不到共同的詞時靠它——
    「手機」跟「行動電話」一個字都不一樣，光看字面永遠對不上。
    """
    template = re.sub(r"\[\d+\]", "[]", key)
    return [label for label, mapped in LABEL_ALIASES.items() if mapped == template]


def _support(slot: Slot, key: str, fields: Dict[str, str]) -> int:
    """這個位置周圍印的字，跟這項資料有幾個詞（連續兩字）相同。

    比對的詞是資料的名稱；一筆一筆的清單資料再加上同一筆的「關係」欄——
    「您將推薦上述哪位主管」跟「諮詢人姓名」沒有共同的詞，靠的是關係欄的「直屬主管」。
    只加關係欄、不加其他值：公司名（「○○公司」）會跟「您對本公司的了解」對上「公司」，
    放行把部門填進問答題；證照機構「電腦技能基金會」會跟「電腦操作」對上「電腦」。
    欄名印英文時（「E-Mail」）中文名稱一個字都對不上，改比代碼本身。
    """
    spec = _spec(key)
    entry = re.match(r"(\w+\[\d+\])\.", key)
    relation = fields.get(f"{entry.group(1)}.relation", "") if entry else ""
    words = _bigrams(spec.label if spec else "") | _bigrams(relation)
    around = _around(slot)
    word = key.split(".")[-1].replace("_", "").lower()
    letters = re.sub(r"[^a-z]", "", around.lower())
    squashed = _squash(around)
    return (len(words & _bigrams(around))
            + (1 if word and word in letters else 0)
            + (2 if any(a in squashed for a in _aliases(key)) else 0)
            + (1 if _period(slot) and _period(slot) == _field_period(key) else 0))


def _period(slot: Slot) -> str:
    """金額後面印的是「/月」還是「/年」：「NT$＿＿/年」要的是年薪。沒有就空字串。"""
    text = slot.cell.paras[slot.para].text
    m = re.match(r"[ 　]*[/／每][ 　]*(月|年)", text[slot.end:])
    return m.group(1) if m else ""


def _field_period(key: str) -> str:
    """金額欄位是月薪還是年薪（看名稱與說明：「期望年薪」「月薪金額」）。認不出就空字串。"""
    spec = _spec(key)
    text = f"{spec.label}{spec.hint}" if spec and spec.kind == "money" else ""
    return "年" if "年薪" in text else "月" if "月薪" in text else ""


# 選項型與長文的資料寫成字時，那一格附近要印著它的名稱（見 _plausible）。表格常用別的
# 說法，只比名稱會擋掉對的配對：「自我介紹」就是自傳、「婚姻：」就是婚姻狀況。
# 只給這道篩選用，不放進 schema.LABEL_ALIASES——那份是讀文字路線的確定性對照，
# 放太鬆會錨錯格。也不改成「有兩個字相同就算」：「狀況」會讓健康狀況放行婚姻狀況
NAME_SYNONYMS = {
    "autobiography": ("自我介紹", "自我簡介", "自述", "個人簡介"),
    "skills.languages": ("語言能力", "外語能力", "語言", "外語"),
    "skills.certificates": ("證照", "證書", "檢定"),
    "skills.computer": ("電腦能力", "電腦專長", "電腦程度", "電腦"),
    "skills.driver_license": ("駕駛執照", "駕駛"),
    "basic.health": ("健康",),
    "basic.marital_status": ("婚姻", "婚否"),
    "basic.military": ("兵役", "役別"),
    "basic.transport": ("交通",),
    "basic.identity_category": ("身分", "身份"),
    "experience[].is_supervisor": ("主管",),
    "declaration.relatives_in_company": ("親友", "親屬"),
    "declaration.other_positions": ("負責人", "董監事"),
    "declaration.china_investment": ("大陸",),
    "declaration.non_compete": ("競業",),
    "declaration.ip_ownership": ("智慧財產", "專門技術"),
    "declaration.criminal_record": ("刑事", "犯罪", "前科"),
    "declaration.wanted": ("通緝",),
    "declaration.infectious_disease": ("傳染病",),
    "declaration.drug_use": ("毒品",),
    "declaration.dismissed": ("免職", "開除", "解僱"),
    "declaration.forged_documents": ("不實",),
    "declaration.debt": ("負債",),
}


def _plausible(slot: Slot, key: str, fields: Dict[str, str]) -> bool:
    """擋掉明顯放錯地方的配對。

    選項型（有／無、是／否）與長文（自傳）的值，離開印著它名稱的地方就沒有意義：
    「否」寫進「直屬主管姓名」、「無」寫進問答題、自傳寫進「本業專長」，都是模型
    拿不相干的資料頂替。所以這兩種資料寫成字時，那一格附近必須印著它的名稱；
    勾選框不受此限（「□否 □是」本來就不會印出資料名稱）。

    一筆一筆的清單資料（學經歷、家人、推薦人）至少要跟那一格的欄名，或同一筆的
    其他資料，有一個詞相同：公司名被放進家庭成員的「姓名」欄時，兩邊一個詞都
    對不上。同一筆的其他資料要算進來，是因為「您將推薦上述哪位主管」這種問句跟
    「諮詢人姓名」沒有共同的詞，靠的是同一筆的關係欄「直屬主管」。
    """
    spec = _spec(key)
    around = _squash(_around(slot))
    value = fields.get(key, "")
    if slot.option == "()":             # 區碼位置只收市話號碼
        return bool(_area_code(slot, value))
    if _period(slot) and _field_period(key) and _period(slot) != _field_period(key):
        return False                    # 月薪不寫進「NT$＿＿/年」
    if slot.kind != "box" and not re.search(r"\d", value):
        # 數字的位置：「每分鐘：＿字」「＿歲」「＿公分」，以及「護照號碼：＿」這種號碼欄。
        # 值裡一個數字都沒有就不是這裡的——英文姓氏被放進「英文輸入 每分鐘：＿字」
        text = slot.cell.paras[slot.para].text
        if (re.match(r"(字|歲|元|公分|公斤|cm|kg)(?![一-鿿])", text[slot.end:].lstrip(" 　"))
                or re.search(r"(號碼|字號|編號)[.．]?[：:]?[ 　]*$", text[:slot.start])):
            return False
    if slot.kind != "box" and spec and spec.kind in ("choice", "longtext"):
        names = (spec.label, *NAME_SYNONYMS.get(re.sub(r"\[\d+\]", "[]", key), ()))
        return any(_squash(n) in around for n in names)
    if re.match(r"\w+\[\d+\]\.", key):
        return _support(slot, key, fields) > 0
    if slot.kind == "box":
        return True
    # 其他單項資料寫成字時，名稱至少要有一個字出現在那一格附近；英文欄名比對代碼
    # （「E-mail」對 contact.email）。擋的是「語言:1.」後面填興趣、填本人姓名這種
    word = key.split(".")[-1].replace("_", "").lower()
    return (bool(set(re.findall(r"[一-鿿]", spec.label if spec else "")) & set(around))
            or word in re.sub(r"[^a-z]", "", around.lower())
            or any(a in around for a in _aliases(key)))


def _as_tick(by_id: Dict[str, Slot], out: Dict[str, str], fields: Dict[str, str]) -> None:
    """值就是框旁邊印的選項時，那一項用勾表達，不是寫在同一格的空位上。

    「負債狀況：▁　□無　□卡債：$」——模型指的那一格沒錯，但這一格填的是勾 ■無，
    在冒號後面再寫一個「無」就多出一個字。同一格裡有字面對得上的框就讓給框。

    只在這一格整個就一題勾選題時才動。宣告事項那種一格印六題的，對得上的框
    很可能是別題的，換過去就勾錯行。
    """
    by_addr: Dict[str, List[Slot]] = {}
    for x in by_id.values():
        by_addr.setdefault(x.addr, []).append(x)
    for sid, key in list(out.items()):
        slot = by_id[sid]
        if slot.kind == "box":
            continue
        boxes = [b for b in by_addr[slot.addr] if b.kind == "box"]
        if len(box_groups(boxes)) != 1:
            continue
        box = next((b for b in boxes if b.id not in out
                    and _ticked(b.option, fields.get(key, ""))), None)
        if box is not None:
            log.info("改用勾的 %s：%s → %s", key, sid, box.id)
            del out[sid]
            out[box.id] = key


def dedupe(slots: List[Slot], chosen: Dict[str, str], fields: Dict[str, str],
           trusted: frozenset = frozenset()) -> Dict[str, str]:
    """先擋掉放錯地方的配對，再讓同一項資料在同一張表格裡只留一個格子：
    搶同一項的格子裡，欄名跟資料名稱最像的那格贏，平手看誰在前面。

    trusted 是「整張表問過欄名」得到的配對，不再用詞比對檢查——模型把「服務單位」
    判成公司名稱是對的，但那兩個詞一個字都不一樣，用詞比對會把整欄擋掉。

    模型的錯多半是「找不到更好的就拿用過的頂替」：學歷表列數比資料多就把第一筆
    往下每列再填一次、住家電話欄填行動電話、業外專長欄再填一次本業專長。
    同一張表格內不重複，這些一次全部擋掉。跨表格不管——姓名出現在表頭又出現在
    簽名欄是正常的。同一格內重複也不管：「□未婚 □已婚」兩個框都指向婚姻狀況。
    """
    by_id = {s.id: s for s in slots}
    out = {sid: key for sid, key in chosen.items()
           if sid in trusted or _plausible(by_id[sid], key, fields)}
    _as_tick(by_id, out, fields)
    order: Dict[str, int] = {}
    claims: Dict[Tuple[str, str], Dict[str, List[Slot]]] = {}
    for i, slot in enumerate(slots):                  # slots 照文件順序
        order.setdefault(slot.addr, i)
        key = out.get(slot.id)
        if key:
            table = slot.addr.split(".")[0]
            claims.setdefault((table, key), {}).setdefault(slot.addr, []).append(slot)
    for (_table, key), by_cell in claims.items():
        # 只留第一個搶到的格子的話，早一步放錯的會把對的偷走：公司名先被放進
        # 家庭成員欄，後面服務單位欄就拿不到了
        winner = max(by_cell, key=lambda a: (any(x.id in trusted for x in by_cell[a]),
                                             max(_support(x, key, fields) for x in by_cell[a]),
                                             -order[a]))
        for addr, group in by_cell.items():
            for slot in (group[1:] if addr == winner else group):
                # 同一格裡的勾選框與空格可以共用同一項資料：「□未婚 □已婚」兩個框都指向
                # 婚姻狀況、「□　8 月 24 日」是勾一個框再把日期寫進後面的空格。
                # 整段值的位置（空格子、冒號後面、問句下一行）就不行——問答格常有
                # 「問句後面」和「空白段」兩個位置，兩個都填會把整段答案寫兩次
                # 沒印單位的兩個空格（「▁　, ▁」英文名與姓氏）不行：值攤不開，第二格只會空著，
                # 讓出來給補漏問
                first_gap = next((x for x in group if x.kind == "gap"), None)
                spreadable = slot.kind == "box" or slot.option in UNITS or slot is first_gap
                if addr == winner and {group[0].kind, slot.kind} <= {"box", "gap"} and spreadable:
                    continue
                del out[slot.id]

    # 同一格內的配對重排：模型偶爾把兩個位置對調（「役別：＿(如免役說明原因)：＿」
    # 填成「癲癇／免役」）。各自的欄名分得出來的話就換回來
    in_cell: Dict[str, List[Slot]] = {}
    for slot in slots:
        if out.get(slot.id) and slot.kind != "box":
            in_cell.setdefault(slot.addr, []).append(slot)
    for addr, group in in_cell.items():
        if not 2 <= len(group) <= 4:
            continue
        keys = [out[x.id] for x in group]
        best = None
        score = sum(_support(x, k, fields) for x, k in zip(group, keys))
        for perm in permutations(keys):
            if perm == tuple(keys):
                continue
            gain = sum(_support(x, k, fields) for x, k in zip(group, perm))
            if gain > score and all(_plausible(x, k, fields) for x, k in zip(group, perm)):
                best, score = perm, gain
        if best:
            log.info("同一格內重排 %s：%s → %s", addr, keys, list(best))
            for x, k in zip(group, best):
                out[x.id] = k

    # 跨表格可以重複（離職原因同時是工作經歷欄與「為何離職」問答題的答案），
    # 但第二張表那裡必須印著跟它有關的字——模型會把本人姓名塞進「語言:1.」後面
    by_key: Dict[str, List[Slot]] = {}
    for slot in slots:
        if out.get(slot.id):
            by_key.setdefault(out[slot.id], []).append(slot)
    for key, group in by_key.items():
        best = max(group, key=lambda x: (x.id in trusted, _support(x, key, fields)))
        for slot in group:
            if (slot.id not in trusted
                    and slot.addr.split(".")[0] != best.addr.split(".")[0]
                    and _support(slot, key, fields) == 0):
                del out[slot.id]
    return out


def _owned(box: Slot, near: List[Slot]) -> List[Slot]:
    """這個框管著的空位：同一段同一行、在它後面、中間沒有別的框的那些非框位置
    （「□ NT$＿＿/月」的空格、「□是，請說明：＿＿」的空格）。"""
    def line(x: Slot) -> int:
        return x.cell.paras[x.para].text.count("\n", 0, x.start)
    same = [x for x in near if x.para == box.para and line(x) == line(box)]
    after = sorted((x for x in same if x.start > box.start), key=lambda x: x.start)
    out = []
    for x in after:
        if x.kind == "box":
            break
        out.append(x)
    return out


def apply_fills(slots: List[Slot], chosen: Dict[str, str], fields: Dict[str, str],
                ticks: Dict[str, bool] = None, highlight: bool = False) -> int:
    """把選中的值寫回文件。ticks 是 ask_boxes 判斷過意思的勾選框，有給就照它勾。
    highlight 把填進去的字標成黃底，只給網頁預覽用。"""
    ticks = ticks or {}
    texts = {x.id: x.cell.paras[x.para].text for x in slots}
    by_id = {x.id: x for x in slots}
    typed_ids = {sid for sid, key in chosen.items() if key and key.startswith(TYPED)}
    reps: Dict[str, str] = {}
    by_cell: Dict[str, List[Slot]] = {}
    by_para: Dict[Tuple[str, int], List[Slot]] = {}
    for slot in slots:
        by_cell.setdefault(slot.addr, []).append(slot)
        by_para.setdefault((slot.addr, slot.para), []).append(slot)

    for _addr, group in by_cell.items():
        # 先處理勾選框：選項字面對得上值就勾；對不上的照 ask_boxes 判斷的意思勾
        for slot in group:
            value = fields.get(chosen.get(slot.id) or "", "")
            if slot.kind == "box" and slot.id in ticks:
                reps[slot.id] = CHECKED if ticks[slot.id] else ""
            elif value and slot.kind == "box":
                reps[slot.id] = _replacement(slot, value)
            elif slot.id in ticks:      # 欄名是選項的空格子（區部底下的「日間」）：打記號，不寫值
                reps[slot.id] = MARK if ticks[slot.id] else ""
        # 已經用打勾表達過的資料，不要在同一格再寫一次字：「■良好」勾好了，
        # 後面「請說明原因：」就不該再補一個「良好」。框後面自己帶著空格、選項字又不是
        # 那個值的（「□ NT$＿/月」「□　＿月＿日後」）不算表達過——值要寫在它的空格裡
        ticked = {chosen.get(x.id) for x in group if x.kind == "box" and reps.get(x.id)
                  and (_ticked(x.option, fields.get(chosen.get(x.id) or "", ""))
                       or not _owned(x, group))} - {None}
        for slot in group:
            key = chosen.get(slot.id)
            value = fields.get(key or "")
            if (value and slot.kind not in ("gap", "box") and key not in ticked
                    and slot.id not in ticks):
                reps[slot.id] = _replacement(slot, value)

        # 空格以「整格」為單位：被指名的那一格往後看，把同一項資料攤到連續的空格上。
        # 跨段落也要算——「自　年　月」與「至　年　月」是同一格裡的兩段，
        # 卻是同一個任職期間
        gaps = [x for x in group if x.kind == "gap" and chosen.get(x.id) not in ticked]
        i = 0
        while i < len(gaps):
            key = chosen.get(gaps[i].id)
            if not fields.get(key or ""):
                i += 1
                continue
            run = [gaps[i]]
            while (not key.startswith(TYPED)      # 自己打的字只寫這一格，不往後面攤
                   and i + len(run) < len(gaps)
                   and chosen.get(gaps[i + len(run)].id) in (None, key)):
                run.append(gaps[i + len(run)])
            reps.update(_spread(run, fields[key], texts))
            i += len(run)
        # 日期單位前面的空格只收數字：「自＿年＿月」塞進一個公司名一定是配錯了
        for gap in gaps:
            if (gap.option in DATE_UNITS and reps.get(gap.id)
                    and not re.fullmatch(r"[\d.,]+", reps[gap.id].strip())
                    and not PRESENT_RE.match(reps[gap.id])):     # 「至＿年」寫「至今」可以
                del reps[gap.id]

    # 一格一個字的格子（身分證字號底下十格）：字數剛好對上才一格一個寫進去。
    # 對不上就整排不寫——整串塞進第一個小格子比留白還糟，而且字數不合多半是配錯了
    for leader, run in char_runs(slots).items():
        if leader in typed_ids:        # 手打的照原樣寫在那一格，不拆成一格一個字
            continue
        chars = re.sub(r"[\s\-－]", "", fields.get(chosen.get(leader) or "", ""))
        reps.pop(leader, None)
        if len(chars) == len(run):
            reps.update({x.id: ch for x, ch in zip(run, chars)})

    for (_addr, _pi), group in by_para.items():
        text = texts[group[0].id]
        for line in {text.count("\n", 0, x.start) for x in group}:
            here = [x for x in group if text.count("\n", 0, x.start) == line]
            boxes = sorted((x for x in here if x.kind == "box"), key=lambda x: x.start)

            def owner(x: Slot):         # 管著這個位置的框：同一行、在它前面最近的那一個
                return next((b for b in reversed(boxes) if b.start < x.start), None)

            # 選項印的字一模一樣、後面各自帶著空格的框（「□ NT$＿＿/月 □ NT$＿＿/年」），
            # 字面分不出該勾哪個，看的是框後面的空格有沒有寫東西：月薪那格寫了就勾月薪那個框。
            # 沒帶空格的（「1.說□很好 □尚可 2.寫□很好 □尚可」）照原本的判斷
            for b in boxes:
                if (sum(_squash(o.option) == _squash(b.option) for o in boxes) > 1
                        and _owned(b, here)):
                    reps[b.id] = CHECKED if any(reps.get(x.id) for x in _owned(b, here)) else ""
            if not any(x.kind == "box" and reps.get(x.id) for x in here):
                for filled in [x for x in here if x.kind in ("gap", "append") and reps.get(x.id)
                               and _kind(chosen.get(x.id) or "") == "date"]:
                    # 勾離那個空格最近的前一個框：「□隨時 □　週 □　8 月 24 日」
                    # 要勾的是第三個。不管那個框被指到哪一項資料——選項本身沒有字
                    # （後面接的是空格）時，模型指到誰都一樣，位置才是意思所在
                    near = [x for x in here if x.kind == "box" and x.start < filled.start
                            and x.id not in ticks]
                    if near:
                        reps[max(near, key=lambda x: x.start).id] = CHECKED
                        break
            # 勾選題後面接的「請說明：」「姓名及部門：」「其他：」是補充說明，只有勾了
            # 前面的選項才寫：這一行勾的是「否／無／不…」、或整行一個都沒勾，就不寫字
            if not boxes:
                continue
            picked = [x for x in boxes if reps.get(x.id)]
            if picked and not any(re.fullmatch(r"否|無|沒有|不.*", x.option) for x in picked):
                # 勾的是別的選項時，沒勾的那個選項後面的空格也不寫：「■台灣 □其他＿＿」
                for x in here:
                    if x.kind != "box" and owner(x) and not reps.get(owner(x).id):
                        reps.pop(x.id, None)
                continue
            for x in here:
                if x.kind != "box" and x.start > boxes[0].start:
                    reps.pop(x.id, None)

    # 使用者自己打的字照原樣寫。上面那些規則（日期拆進年月日、同一項攤到連續空格、
    # 一格一個字、單位前只收數字）都是為了「從我的資料推出這一格該寫什麼」，
    # 手打的不必再推一次——他要什麼就寫什麼。勾選框例外：打字的意思是勾那個選項
    for sid in typed_ids:
        slot = by_id.get(sid)
        if slot is not None:
            text = fields.get(chosen[sid], "")
            reps[sid] = _replacement(slot, text) if slot.kind == "box" else text

    written = 0
    for (_addr, pi), group in by_para.items():
        changes = [(x.start, x.end, reps[x.id]) for x in group if reps.get(x.id)]
        if not changes:
            continue
        for slot in group:
            if reps.get(slot.id):
                written += 1
                log.debug("寫 %-18s %-26s %d字", slot.id, chosen.get(slot.id),
                          len(reps[slot.id]))
        para = group[0].cell.paras[pi]
        write_changes(para, changes, highlight)
        # 位置彼此不重疊，寫完的字應該正好是「原本的字依序換掉這幾段」。
        # 對不上表示位置對回 run 時算錯了——只記格子代碼，不記值
        text, expect, cursor = texts[group[0].id], [], 0
        for start, end, rep in sorted(changes):
            expect += [text[cursor:start], rep]
            cursor = end
        if para.text != "".join(expect) + text[cursor:]:
            log.warning("段落寫完的字跟預期不同 %s", group[0].id)
    return written


ROWS_SYSTEM = """表格裡有幾張「一列填一筆」的表（學歷、工作經歷、家庭成員這類），
下面列出每一張表的欄名。請判斷每一欄要填清單資料的哪一個子欄位。

規則：
1. 只能從給定的代碼挑；那一欄在個人資料裡沒有對應，填 __無__。
2. 看欄名的意思判斷：「服務單位」是公司名稱、「工作期間」是任職期間。
3. 同一張表左右印了兩組一樣的欄名（家庭成員常見），兩組都照樣對應。"""
# 試過加一條「欄名寫成『區部／日間』的是選項，整組都填那個子欄位」：真建築工作經歷表的
# 第一欄（服務單位公司行號）整欄變成答無、AT-1 推薦人的姓名欄變成關係。9B 模型對這段
# 提示很敏感，選項欄改由程式補齊（見 ask_rows 裡的「兄弟欄」），不寫進提示


def _col(slot: Slot) -> int:
    return int(slot.addr.split(".")[2][1:])


def char_runs(slots: List[Slot]) -> Dict[str, List[Slot]]:
    """一格寫一個字的格子：同一列裡連著四格以上的空格子，頭上是同一個欄名
    （「身分證字號」底下切成十格）。回傳 {第一格的位置編號: 整排位置}。

    模型只問第一格；寫的時候字數剛好等於格數，才一格一個字攤開。
    """
    per_cell = Counter(x.addr for x in slots)
    rows: Dict[Tuple[str, str, int], List[Slot]] = {}
    for x in slots:
        if (x.kind == "blank" and x.addr.startswith("t") and per_cell[x.addr] == 1
                and x.cell.col_head and x.cell.head_row >= 0):
            table, row = x.addr.split(".")[:2]
            rows.setdefault((f"{table}.{row}", x.cell.col_head, x.cell.head_row), []).append(x)
    out = {}
    for run in rows.values():
        run.sort(key=_col)
        if len(run) >= 4 and _col(run[-1]) - _col(run[0]) == len(run) - 1:
            out[run[0].id] = run
    return out


def row_blocks(slots: List[Slot]) -> Dict[Tuple[str, int], List[Slot]]:
    """一列一筆的表：左邊整排沒有欄位名稱（或只有列首）、上面有欄名的格子，依表頭分組。

    每一格挑代表位置：空格子就是它自己；只印格式的格子（「自　年　月／至　年　月」）
    取第一個空格，值由 _spread 往後攤；整格都是勾選框的（「□畢 □肄」）取全部框。
    表頭底下只有一列的不算（緊急聯絡人「姓名｜關係」底下一列）——那不是清單，
    是一個人的幾項資料，照一般的位置問。
    """
    out: Dict[Tuple[str, int], List[Slot]] = {}
    by_cell: Dict[str, List[Slot]] = {}
    in_runs = {x.id for run in char_runs(slots).values() for x in run}
    for x in slots:
        if x.cell.depth and x.id not in in_runs:
            by_cell.setdefault(x.addr, []).append(x)
    for addr, group in by_cell.items():
        kinds = {x.kind for x in group}
        if "blank" in kinds:
            rep = [x for x in group if x.kind == "blank"][:1]
        elif kinds == {"box"}:
            rep = group
        elif "gap" in kinds:
            rep = [x for x in group if x.kind == "gap"][:1]
        else:
            continue
        out.setdefault((addr.split(".")[0], group[0].cell.origin), []).extend(rep)
    return {k: group for k, group in out.items()
            if len({x.cell.depth for x in group}) >= 2}


def _names_row(value: str, key: str) -> bool:
    """列首印的就是這一筆的值：「大學」對「大 學」、「專科」對「專 科(二.三.五專)」、
    「碩士」對「碩 / 博士」。值的每個中文字都要出現在列首，而且至少兩個字——
    「日」「畢」這種一個字的值什麼列首都沾得上。"""
    chars = re.findall(r"[一-鿿]", value or "")
    return len(chars) >= 2 and all(ch in key for ch in chars)


def _option_hit(value: str, head: str) -> bool:
    """選項欄的欄名跟資料的值是同一個選項：「日」對「日間」、「畢」對「畢業」、「日間部」對「日間」。"""
    v, h = _squash(value), _squash(head)
    return bool(v and h) and (h.startswith(v) or v.startswith(h))


def _row_records(group: List[Slot], root: str, fields: Dict[str, str]) -> Dict[int, int]:
    """列首印著這一筆是哪一種的表（碩／博士｜大學｜專科｜高中(職)）：{第幾列: 第幾筆}。

    找出「哪一個子欄位的值印在列首」（學歷的學位），再照值把每一筆放到對得上的
    那一列。一列對到兩筆、或一筆對到兩列，那個子欄位就不算數——分不清就不放。
    對不上任何一列的資料不填：表格沒有那一類的列，硬塞進別列一定是錯的。
    """
    rows = {x.cell.depth: x.cell.key for x in group if x.cell.key}
    count = len({re.match(r"\w+\[(\d+)\]", k).group(1) for k in fields if k.startswith(root + "[")})
    subs = {k.split("].", 1)[1] for k in fields if k.startswith(root + "[")}
    best: Dict[int, int] = {}
    for sub in sorted(subs):
        hits = {d: [i for i in range(count)
                    if _names_row(fields.get(f"{root}[{i}].{sub}", ""), key)]
                for d, key in rows.items()}
        hits = {d: ids for d, ids in hits.items() if ids}
        records = [i for ids in hits.values() for i in ids]
        if (all(len(ids) == 1 for ids in hits.values()) and len(records) == len(set(records))
                and len(hits) > len(best)):
            best = {d: ids[0] for d, ids in hits.items()}
    return best


def ask_rows(blocks: Dict[Tuple[str, int], List[Slot]], fields: Dict[str, str],
             pages: Pages, printed: str, host: str = LLM_HOST,
             model: str = LLM_MODEL) -> Tuple[Dict[str, str], Dict[str, bool]]:
    """一列一筆的表：問模型「每一欄是什麼」，再由程式照「往下第 N 列＝第 N 筆」排進去。

    逐格問時，這種表的每一格都長得一模一樣（空的，只差欄名），模型會把第二筆
    錯位到第三列、或整列漏掉第三筆。但「服務單位公司行號這一欄是公司名稱」
    它答得很穩——欄的意思交給模型，列的順序交給程式。

    第幾筆放哪一列，依序看三件事：列首印著這一筆是哪一種（大學那筆放「大學」列）、
    左邊序號欄印的數字（證照表左欄 1、2，右欄 3、4）、都沒有才照由上而下的順序。
    回傳 {位置編號: 項目代碼} 與 {位置編號: 打不打記號}（選項欄用）。

    printed＝這份表格印出來的字，沒有預設值是故意的：少傳了就等於「什麼都沒提到」，
    家人、諮詢人那幾份清單會整個消失，而且不會有任何錯誤訊息。
    """
    roots = {k.split("[")[0] for k in fields if "[" in k}
    choices = [k for k in BY_KEY if "[]" in k and k.split("[]")[0] in roots
               and _mentioned(k, printed) and _worth_offering(k, fields)]
    ids, lines, allowed = [], [], {}
    for bi, group in enumerate(blocks.values(), 1):
        # 左邊的區塊標題；序號（「1」）與列首（「大學」）不是在說這一區是什麼
        section = (next((x.cell.row_head for x in group if x.cell.row_head and not x.cell.key
                         and not _is_template(x.cell.row_head)), "")
                   or next((x.cell.caption for x in group if x.cell.caption), ""))
        # 這一區標了是誰（「家庭成員」、「請列舉…並同意我們諮詢」），可選的就只剩
        # 那一組。不先收斂的話，欄名「稱謂／姓名／服務機關／職位」跟諮詢人太像，
        # 整張表會被對到另一份名單。比對前要抹空白：標題常寫成「家　庭　成　員」
        picks = [k for root, words in OTHER_PEOPLE.items()
                 if any(w in _squash(section) for w in words)
                 for k in choices if k.startswith(root)]
        keys = list(dict.fromkeys(_squash(x.cell.key) for x in group if x.cell.key))
        lines.append(f"表{bi}" + (f"（這一區印著「{_squash(section)[:40]}」）" if section else "")
                     + (f"（每一列的列首：{'、'.join(keys)}）" if keys else "") + "：")
        for c, head, parent in sorted({(_col(x), x.cell.col_head, x.cell.col_parent)
                                       for x in group}):
            ids.append(f"b{bi}c{c}")
            allowed[f"b{bi}c{c}"] = picks or choices
            lines.append(f"  b{bi}c{c}｜欄名：" + (f"{_squash(parent)}／" if parent else "")
                         + _squash(head))
    schema = {
        "type": "object",
        "properties": {i: {"type": "string", "enum": allowed[i] + [NONE]} for i in ids},
        "required": ids,
        "additionalProperties": False,
    }
    user: List[Dict[str, Any]] = [{"type": "text", "text": (
        "可用的代碼（附第一筆的值，讓你對得上欄名）：\n"
        + "\n".join(f"- {k}{_label(k)}"
                    + (f"：{fields[k.replace('[]', '[0]')][:20]}"
                       if k.replace("[]", "[0]") in fields else "")
                    for k in choices)
        + "\n\n各表的欄名：\n" + "\n".join(lines))}]
    # 欄名印在資料列上面那一列，那一列也要看得到——表格跨頁時欄名與資料會分在兩張上
    addrs = {x.addr for group in blocks.values() for x in group}
    addrs |= {a for table, head_row in blocks
              for a in pages.page_of if a.startswith(f"{table}.r{head_row}.c")}
    user += pages.parts(pages.for_addrs(addrs))
    data = llm.ask(host, ROWS_SYSTEM, user, schema, model=model,
                   label=f"一列一筆:{len(blocks)}表")

    out: Dict[str, str] = {}
    marks: Dict[str, bool] = {}
    for bi, group in enumerate(blocks.values(), 1):
        mapping = {_col(x): data.get(f"b{bi}c{_col(x)}") for x in group}
        mapping = {c: k for c, k in mapping.items() if k in choices}
        # 一張表就是一份清單。欄位分屬兩份清單時，少數派的那幾欄是模型串了行——
        # 工作經歷表的「直屬主管姓名」問的是這一列的主管，不是諮詢人那一份名單
        main = ""
        if mapping:
            main = Counter(k.split("[")[0] for k in mapping.values()).most_common(1)[0][0]
            mapping = {c: k for c, k in mapping.items() if k.split("[")[0] == main}
        # 同一層表頭底下好幾個欄對到同一個子欄位（區部底下的日間｜夜間｜假日都是
        # 日夜間部）：那幾欄是選項，值對得上欄名的那一欄打記號，其他欄空著。
        # 要有共同的上層表頭才算——模型把「服務機關」「職位」都對到職業時，那是對錯了，
        # 不是選項
        heads_of = {_col(x): x.cell.col_head for x in group}
        parent_of = {_col(x): x.cell.col_parent for x in group}
        # 模型常只把其中一欄（畢業）對到子欄位，另外兩欄（肄業、在學）答無。同一層表頭
        # 底下的兄弟欄都沒對到別的東西、而且資料的值就是其中一欄的欄名時，整組都算
        for c, k in list(mapping.items()):
            siblings = [o for o in heads_of if parent_of[o] and parent_of[o] == parent_of[c]]
            values = [v for f, v in fields.items() if re.sub(r"\[\d+\]", "[]", f) == k]
            if (len(siblings) > 1 and all(mapping.get(o) in (None, k) for o in siblings)
                    and any(_option_hit(v, heads_of[o]) for v in values for o in siblings)):
                mapping.update({o: k for o in siblings})
        options = {c for c, k in mapping.items() if parent_of[c] and len(
            {_squash(heads_of[o]) for o, k2 in mapping.items()
             if k2 == k and parent_of[o] == parent_of[c]}) > 1}
        placed = _row_records(group, main, fields) if main else {}
        numbered = all(x.cell.seq for x in group)
        # 一列放幾筆，看表格自己印了幾組同樣的欄名（家庭成員左右各一組「姓名 稱謂
        # 年齡 職業」），由左而右、由上而下編號。不能看模型的答案——模型只要把兩欄
        # 對到同一個子欄位，整張表就被當成一列兩筆，第二列拿到第三筆
        heads = {c: _squash(h) for c, h in {(_col(x), x.cell.col_head) for x in group}}
        per_row = max(Counter(heads.values()).values(), default=1)
        seen: Counter = Counter()
        nth = {}
        for c in sorted(heads):
            nth[c] = seen[heads[c]]
            seen[heads[c]] += 1
        log.info("一列一筆 表%d：%s%s", bi, {c: mapping[c] for c in sorted(mapping)},
                 f"（列首對應：{placed}）" if placed else "")
        for x in group:
            c = _col(x)
            if c not in mapping:
                continue
            if placed:
                if x.cell.depth not in placed:
                    continue
                index = placed[x.cell.depth]
            elif numbered:
                index = x.cell.seq - 1
            else:
                index = (x.cell.depth - 1) * per_row + nth[c]
            key = mapping[c].replace("[]", f"[{index}]")
            if key not in fields:
                continue
            if c in options:
                if x.kind == "blank" and _option_hit(fields[key], heads_of[c]):
                    out[x.id], marks[x.id] = key, True
                continue
            out[x.id] = key
    return out, marks


RECALL_SYSTEM = """表格上還有幾個位置空著，個人資料裡也還有沒填進去的項目。
每個位置後面列了幾項名稱相近、還沒用到的資料。每個位置先挑一項（pick），
再判斷挑的這一項是不是真的就是那一格要的東西（answers）。

規則：
1. pick：只能從那個位置列出的項目裡挑；都不對就填 __無__。
2. answers：挑的資料就是那一格問的東西才填「是」，只是字面有點像不算。
   例如位置問「目前住址」、資料是「緊急聯絡人的住址」，要填「否」。"""


def _placed(slots: List[Slot], chosen: Dict[str, str],
            fields: Dict[str, str]) -> set:
    """實際會寫出東西的資料項。

    「被指派過」不等於「填進去了」：日期被指到「□隨時」那個框，勾不起來、也沒寫進
    任何空格，等於沒填，要留給補漏。選項型的資料（婚姻、健康、宣告事項）被指到
    勾選框就算表達完了——字面對不上（資料「無」、表格印「□否」）由 ask_boxes 判斷。
    """
    out = set()
    for s in slots:
        key = chosen.get(s.id)
        value = fields.get(key or "")
        if value and (s.kind != "box" or _ticked(s.option, value) or _kind(key) == "choice"):
            out.add(key)
    return out


def recall(slots: List[Slot], chosen: Dict[str, str], fields: Dict[str, str],
           pages: Pages, skip: frozenset = frozenset(),
           host: str = LLM_HOST, model: str = LLM_MODEL) -> Dict[str, str]:
    """補漏：還空著的位置，只拿「名稱相近、還沒用到」的資料再問一次。

    第一輪一百多個位置逐個問，模型對某一格想不出來就答 __無__，漏掉的不會再有人問。
    反過來問「這幾十項資料各該去哪」也試過：9B 模型把五十幾項全塞進同一個位置。
    改成每個空位置只列出跟它有共同詞的幾項未用資料——選項少、每一項都跟那一格
    沾得上邊，是小得多的一道題。
    """
    placed = _placed(slots, chosen, fields)
    left = [k for k in fields if k not in placed]
    todo = []
    for s in slots:
        # 看單一位置還空不空就好：同一行勾好了「■免役」，後面的「原因＿＿」還是要補
        if s.id in chosen or s.id in skip:
            continue
        # 勾選框：資料的值就是印在框旁邊的選項（「畢」對「□畢」），或至少有共同的詞
        # （「104 人力銀行」對「□人力網站」），是比欄名更強的線索
        ranked = sorted(((_support(s, k, fields)
                          + (2 if s.kind == "box" and (_ticked(s.option, fields[k])
                             or _bigrams(fields[k]) & _bigrams(s.option)) else 0), k)
                         for k in left), reverse=True)
        cands = [k for score, k in ranked if score > 0 and _plausible(s, k, fields)][:6]
        if cands:
            todo.append((s, cands))
    if not todo:
        return {}

    ids = [f"s{i}" for i in range(1, len(todo) + 1)]
    # 先挑再判斷「是不是真的在回答那一格」：只挑的話，模型會因為兩邊都有「公司」
    # 就把推薦人的公司填進「您對本公司的了解」——字面對得上，問的卻不是同一家公司
    schema = {
        "type": "object",
        "properties": {i: {
            "type": "object",
            "properties": {"pick": {"type": "string", "enum": cands + [NONE]},
                           "answers": {"type": "string", "enum": ["是", "否"]}},
            "required": ["pick", "answers"],
            "additionalProperties": False,
        } for i, (_s, cands) in zip(ids, todo)},
        "required": ids,
        "additionalProperties": False,
    }
    lines = [f"  {i}｜{_describe(s)}\n      可選：" + "、".join(
        f"{k}{_label(k)}＝{fields[k][:20]}" for k in cands) for i, (s, cands) in zip(ids, todo)]
    user: List[Dict[str, Any]] = [
        {"type": "text", "text": "空著的位置與可選的資料：\n" + "\n".join(lines)}]
    user += pages.parts(pages.for_addrs({s.addr for s, _cands in todo}))
    data = llm.ask(host, RECALL_SYSTEM, user, schema, model=model,
                   label=f"補漏:{len(ids)}處")
    out = {}
    for i, (s, cands) in zip(ids, todo):
        ans = data.get(i) or {}
        if ans.get("pick") not in cands:
            continue
        if ans.get("answers") != "是":
            log.info("補漏不採用 %s → %s（模型判斷不是那一格要的）", ans["pick"], s.id)
            continue
        out[s.id] = ans["pick"]
        log.info("補漏 %s → %s", ans["pick"], s.id)
    return out


BOXES_SYSTEM = """表格上的勾選題。每一題給你題目印的字和它的選項，請回答兩件事：
這一題在問個人資料的哪一項（field），以及照那一項的值該勾哪一個選項（pick）。

規則：
1. field 只能從項目代碼裡挑。個人資料沒有對應的、或這題是公司自己要填的，填 __無__。
2. pick 只能從那一題印出來的選項裡挑，意思相同就算：資料「無」對選項「否」、
   「中等」對「尚可」、「104 人力銀行」對「人力網站」、「普通重型機車」對「機車」。
   意思都不符就填 __無__。
3. 題目跟資料講的不是同一件事，兩個都填 __無__——題目問「可配合加班」而資料是
   「擔任主管」，選項雖然都是是非，卻不是同一件事。
4. 可以複選的題目（電腦操作、交通工具）挑一個就好，其餘對得上的由程式補。"""


def box_groups(slots: List[Slot]) -> Dict[Tuple[str, int, int], List[Slot]]:
    """勾選題分組：同一格、同一段、同一行的框算一題（宣告事項一格印七題）。

    選項太多換到下一行的（「資訊來源：□公司詢問 □人力網站 □員工介紹↵□其他」）還是同一題：
    一行一開頭就是框、上一行也有框，就併進上一行那一題。不併的話「□其他」自己成了
    只有一個選項的題目，模型看得出在問資訊來源，就把它勾起來。
    """
    out: Dict[Tuple[str, int, int], List[Slot]] = {}
    alias: Dict[Tuple[str, int, int], Tuple[str, int, int]] = {}
    last: Dict[str, Tuple[int, Tuple[str, int, int]]] = {}   # 格子 -> (上一個有框的行, 那一題)
    for x in slots:
        if x.kind != "box":
            continue
        text = x.cell.paras[x.para].text
        line = text.count("\n", 0, x.start)
        own = (x.addr, x.para, line)
        if own not in alias:
            at = sum(p.text.count("\n") + 1 for p in x.cell.paras[:x.para]) + line
            prev = last.get(x.addr)
            leading = re.match(f"[ 　]*[{CHECKBOX_CHARS}]", text.split("\n")[line])
            alias[own] = prev[1] if prev and prev[0] == at - 1 and leading else own
            last[x.addr] = (at, alias[own])
        out.setdefault(alias[own], []).append(x)
    return out


def ask_boxes(groups: Dict[Tuple[str, int, int], List[Slot]], fields: Dict[str, str],
              pages: Pages, chosen: Dict[str, str], narrow: bool = False,
              host: str = LLM_HOST,
              model: str = LLM_MODEL) -> Tuple[Dict[str, str], Dict[str, bool]]:
    """勾選題自己問一輪：這一題在問哪一項資料、該勾哪一個選項。

    勾選題跟填空題是兩種題目。夾在一百多個位置裡逐格問時，模型要順便判斷「□尚可」
    在問什麼，常常答不出來；而且表格印的字跟資料的寫法往往不同（「中等」對「尚可」、
    「104 人力銀行」對「人力網站」），那是語意問題，要把題目、選項、資料擺在一起
    才判得準。回傳 {位置編號: 資料代碼} 與 {位置編號: 勾不勾}。
    """
    items = list(groups.items())
    if len(items) > BOX_BATCH:      # 一次問太多題，模型會整批答 __無__
        # 平均分攤，不要留下落單的一題：17 題分成 8+8+1 時，最後那一批只有一題、
        # 前後都沒有同類的題目可以參照，實測整批答 __無__。分成 6+6+5 就好了
        size = -(-len(items) // -(-len(items) // BOX_BATCH))
        picked, ticks = {}, {}
        for i in range(0, len(items), size):
            part = dict(items[i:i + size])
            got, tk = ask_boxes(part, fields, pages, {**chosen, **picked}, narrow,
                                host, model)
            picked.update(got)
            ticks.update(tk)
        return picked, ticks
    ids = [f"q{i}" for i in range(1, len(groups) + 1)]
    options, lines, picks = [], [], []
    for qid, ((_addr, pi, line), boxes) in zip(ids, groups.items()):
        printed = boxes[0].cell.paras[pi].text.split("\n")[line]
        # 這一行一開頭就是框（「是否有配偶…任職？↵ □否 ↵□是，請說明：」）：題目印在
        # 前面那一行，只給「□否」模型不知道在問什麼
        if re.match(f"[ 　]*[{CHECKBOX_CHARS}]", printed):
            cell_lines = "\n".join(p.text for p in boxes[0].cell.paras).split("\n")
            at = sum(p.text.count("\n") + 1 for p in boxes[0].cell.paras[:pi]) + line
            question = next((t for t in reversed(cell_lines[:at]) if t.strip()
                             and not re.match(f"[ 　]*[{CHECKBOX_CHARS}]", t)), "")
            printed = _squash(question)[-36:] + " " + printed
        # 併進來的下一行選項（「↵□其他」）也印出來
        for b in boxes:
            whole = b.cell.paras[b.para].text
            more = whole.split("\n")[whole.count("\n", 0, b.start)]
            if more not in printed:
                printed += " " + more
        opts = list(dict.fromkeys(b.option for b in boxes if b.option))
        options.append(opts)
        ctx = _squash(boxes[0].cell.row_head + boxes[0].cell.col_head)[:20]
        # 附上幾個可能相關的資料：欄名對得上、值就是印在框旁邊的選項、或同一列已經
        # 填過同一群的資料（「語言:1.英文」旁邊那題八成就是問語文程度）。
        # 一百多個項目裡要模型自己撈，實測駕照、負債狀況這種明明對得上的都會漏掉
        row = boxes[0].addr.rsplit(".", 1)[0]
        near = {chosen[k].split(".")[0] for k in chosen if k.rsplit(".", 1)[0].startswith(row)}
        # 排序分三層，不相加：名稱與值對得上是實證，另外兩個只在實證平手時才作數。
        # 相加的話「已經填在別處、名字又剛好沾得上」的項目會靠同列加分擠掉正解——
        # 真建築「您是否有親友在本公司服務？」後面印著「姓名及部門」，basic.name_zh
        # 跟 basic.name_en 就是這樣佔掉兩個名額，declaration.relatives_in_company 掉出前四
        ranked = sorted(((_support(boxes[0], k, fields)
                          + (2 if any(_ticked(b.option, v) or _bigrams(v) & _bigrams(b.option)
                                      for b in boxes) else 0),
                          0 if k in chosen.values() else 1,
                          1 if k.split(".")[0] in near else 0, k)
                         for k, v in fields.items()), reverse=True)
        # 同一個欄位只列一次：三筆工作經歷的「擔任主管＝否」長得一模一樣，
        # 不去掉就把四個名額佔滿，真正對得上的那一項擠不進來
        best, seen = [], set()
        for evidence, _fresh, near_hit, k in ranked:
            template = re.sub(r"\[\d+\]", "[]", k)
            if evidence + near_hit > 0 and template not in seen and len(best) < 4:
                seen.add(template)
                best.append(k)
        hint = "、".join(f"{k}{_label(k)}＝{fields[k][:14]}" for k in best)
        picks.append(best)
        log.debug("勾選候選 %s#%s.%s 題「%s」→ %s", boxes[0].addr, pi, line,
                  _squash(printed)[:24], best)
        lines.append(f"  {qid}｜{_squash(printed)[:70]}" + (f"｜欄名:{ctx}" if ctx else "")
                     + (f"\n      可能相關：{hint}" if hint else ""))
    schema = {
        "type": "object",
        "properties": {qid: {
            "type": "object",
            "properties": {"field": {"type": "string",
                                     "enum": (cands if narrow else list(fields)) + [NONE]},
                           "pick": {"type": "string", "enum": opts + [NONE]}},
            "required": ["field", "pick"],
            "additionalProperties": False,
        } for qid, opts, cands in zip(ids, options, picks)},
        "required": ids,
        "additionalProperties": False,
    }
    user: List[Dict[str, Any]] = [{"type": "text", "text": (
        "個人資料（項目代碼：值）：\n"
        + "\n".join(f"- {k}{_label(k)}：{v[:60]}" for k, v in fields.items()))}]
    # 這一輪每一批各問各的，不接續、也不照批次挑頁：問的是「這題在問哪一項資料」，
    # 兩種省法都會讓模型跳題。只附題目所在那一頁時真建築少一題；接在前一批後面問
    # 時換成另外兩題被跳過（十幾題擺在同一串對話裡，跟一次問十幾題是同一個毛病）。
    # 填寫的那幾輪只要找得到位置，才省得下來
    user += pages.parts(pages.all())
    user.append({"type": "text", "text": "要判斷的勾選題：\n" + "\n".join(lines)})
    data = llm.ask(host, BOXES_SYSTEM, user, schema, model=model,
                   label=f"勾選題{'再問' if narrow else ''}:{len(ids)}題")

    picked: Dict[str, str] = {}
    ticks: Dict[str, bool] = {}
    for qid, (_g, boxes) in zip(ids, groups.items()):
        ans = data.get(qid) or {}
        key = ans.get("field")
        value = fields.get(key or "")
        if not value:
            continue
        pick = ans.get("pick")
        if not _names_question(boxes[0], key, fields):
            log.info("勾選題不採用 %s %s（列首印的是別的東西）", boxes[0].addr, key)
            continue
        # 選項裡沒有「否／無」的題目（「□領有身心障礙手冊 □原住民」），框印的都是肯定的
        # 那一面：資料是「無／否」就一個都不勾，模型說勾也不勾
        lone_no = (not any(re.fullmatch(NEGATIVE, b.option) for b in boxes)
                   and re.fullmatch(NEGATIVE, value.strip()))
        # 同義詞對得上的選項（「中華民國」對「台灣」）就是答案，模型挑了別的也不算
        literal = [b for b in boxes if _synonym_hit(b.option, value)]
        for b in boxes:
            picked[b.id] = key
            # 字面對得上的都勾（複選題「Windows、Word、Excel」一次勾三個），
            # 加上模型判斷意思相同的那一個
            ticks[b.id] = not lone_no and (b in literal or (
                not literal and bool(b.option) and b.option == pick))
        # 值與勾了哪個選項都是個資（勾「有」就是答案），只記勾了幾個
        log.info("勾選題 %s %s → 勾 %d／%d 個", boxes[0].addr, key,
                 sum(ticks[b.id] for b in boxes), len(boxes))
    return picked, ticks


def _synonym_hit(option: str, value: str) -> bool:
    """選項跟值字面相同，或是產品同義詞表裡的同一組（「中華民國」「台灣」）。"""
    if _ticked(option, value):
        return True
    return any(_squash(option) in words and _squash(value) in words
               for words in ({_squash(w) for w in group} for group in OPTION_SYNONYMS))


def _names_question(box: Slot, key: str, fields: Dict[str, str]) -> bool:
    """列首印的是「這一題在問哪一個東西」時，挑的資料要跟那個東西有關。

    語言表一列一種語言：「閩南語｜□優 □普通 □略懂」「其他｜□優 □普通 □略懂」。
    語文程度只屬於資料裡那一種語言（英文）的那一列，放到閩南語、其他那一列就錯了；
    模型還會因為「普通重型機車」裡有「普通」兩個字，把駕照勾進「其他」的程度。
    所以列首是個兩三個字、又不是任何欄位名稱的詞時，同一組資料（skills.*）裡要有
    一項的值跟列首有共同的字（「英文」對「英 語」）才採用。列首是欄位名稱
    （「婚姻」「交通工具」）的照舊。
    """
    row = _squash(box.cell.row_head)
    if not box.addr.startswith("t") or not re.fullmatch(r"[一-鿿]{2,4}", row):
        return True
    names = set(LABEL_ALIASES) | {f.label for f in BY_KEY.values()}
    if any(n.startswith(row) or row.startswith(n) for n in names):
        return True
    group = key.rsplit(".", 1)[0]
    return any(set(value) & set(row) for k, value in fields.items()
               if k.rsplit(".", 1)[0] == group and k != key)


def one_question_per_field(groups: Dict[Tuple[str, int, int], List[Slot]], fields: Dict[str, str],
                           picked: Dict[str, str], ticks: Dict[str, bool]) -> None:
    """同一格裡兩題勾選題對到同一項資料時，只留題目跟它最像的那一題。

    富邦「其他」那一格印了兩題：「是否曾至富邦集團任職過？」與「是否有配偶或二親等血親
    姻親於富邦集團任職？」。親友任職只回答第二題，模型卻兩題都拿它勾「否」——第一題
    app.db 根本沒有資料。兩題的「任職」「公司」一樣多，第二題多了「血親」「姻親」這兩個
    別名，分數高的留下。分數一樣就都留（兩題題目都沒比對不出差別時，不替模型決定）。
    """
    claims: Dict[Tuple[str, str], List[List[Slot]]] = {}
    for boxes in groups.values():
        key = next((picked[b.id] for b in boxes if b.id in picked), None)
        if key:
            claims.setdefault((boxes[0].addr, key), []).append(boxes)
    for (_addr, key), questions in claims.items():
        if len(questions) < 2:
            continue
        score = {id(q): _support(q[0], key, fields) for q in questions}
        best = max(score.values())
        for q in questions:
            if score[id(q)] < best:
                log.info("勾選題不採用 %s %s（同一格另一題更像）", q[0].id, key)
                for b in q:
                    picked.pop(b.id, None)
                    ticks.pop(b.id, None)


def obvious_boxes(groups: Dict[Tuple[str, int, int], List[Slot]], fields: Dict[str, str],
                  picked: Dict[str, str], ticks: Dict[str, bool]) -> Dict[str, str]:
    """模型沒答出來的勾選題，程式自己認得出來的就自己填。

    認的條件是兩邊都成立：題目旁邊印的就是某一項資料的名稱（產品對照表裡的別名，
    「資訊來源」就是招募管道），而且那一項的值跟其中一個選項有共同的詞（「104 人力
    銀行」對「人力網站」）。兩個條件湊在一起已經不是猜——名稱指定了問的是哪一項，
    值指定了該勾哪一個。缺一個就放著不填：「□其他」旁邊也印著「健康狀況」，
    可是健康狀況的值「優」對不上「其他」，那就不是它。

    名稱有好幾個對得上時取離這個位置最近的：一列印兩組欄位是常態，「資訊來源」
    那一格的左邊就是「健康狀況：□優□良」，兩個名稱都在周圍的字裡。
    """
    out: Dict[str, str] = {}
    for boxes in groups.values():
        if any(b.id in picked for b in boxes):
            continue
        around = _squash(_around(boxes[0]))

        def near(key: str) -> int:      # 名稱印在哪裡，越靠近這個位置越可能是它
            spec = _spec(key)
            return max(around.rfind(a)
                       for a in _aliases(key) + [spec.label if spec else ""] if a)

        named = sorted(((pos, k) for pos, k in ((near(k), k) for k in fields)
                        if pos >= 0), reverse=True)
        if not named or (len(named) > 1 and named[0][0] == named[1][0]):
            continue
        key = named[0][1]
        hit = {b.id: _synonym_hit(b.option, fields[key])
               or bool(_bigrams(b.option) & _bigrams(fields[key])) for b in boxes}
        if not any(hit.values()):
            continue
        for b in boxes:
            out[b.id] = key
            ticks[b.id] = hit[b.id]
        # 值與勾了哪個選項都是個資，只記勾了幾個
        log.info("欄名就是它 %s %s → 勾 %d／%d 個", boxes[0].addr, key,
                 sum(hit.values()), len(boxes))
    return out


def _batches(slots: List[Slot]) -> List[List[Slot]]:
    """分批時同一個格子的位置不拆開——「優點：」「缺點：」被拆到兩批問，
    模型看不到彼此就容易錯位。"""
    out: List[List[Slot]] = []
    for _addr, group in groupby(slots, key=lambda s: s.addr):
        group = list(group)
        if out and len(out[-1]) + len(group) <= BATCH:
            out[-1] += group
        else:
            out.append(group)
    return out


@dataclass
class Draft:
    """一份表格分析完的結果。

    slots 不存進資料庫——它牽著 python-docx 的段落物件。要用的時候照原檔重新
    解析一次就好：解析純粹是程式，同一份文件跑幾次結果都一樣，約 0.2 秒。
    存下來的只有 assignment 與 ticks，那兩個才是模型（和使用者）的決定。
    """
    slots: List[Slot]
    fields: Dict[str, str]          # 這份表格用得到的個人資料
    assignment: Dict[str, str]      # 位置編號 -> 項目代碼
    ticks: Dict[str, bool]          # 位置編號 -> 勾不勾


def parse(blank: Path) -> Tuple[Any, List[Cell], List[Slot]]:
    """把文件拆成格子與可寫位置。沒有模型。"""
    doc = Document(str(blank))
    form = cells(doc)
    return doc, form, [s for cell in form for s in slots_of(cell)]


def printed_text(form: List[Cell]) -> str:
    """這份表格上印著的所有字（欄名、說明、選項）。"""
    return "".join(p.text for cell in form for p in cell.paras)


def usable_fields(form: List[Cell], profile: Dict[str, Any]) -> Dict[str, str]:
    """這份表格用得到的個人資料。

    別人的資料（緊急聯絡人、家人、諮詢人）只在表格提到那個人時才拿出來給模型挑：
    沒提到還留在清單裡，模型會把緊急聯絡人的電話填進本人的住家電話欄。
    比對前抹掉空白——標題常寫成「家　庭　成　員」，不抹就對不上「家庭」。
    """
    printed = printed_text(form)
    return {k: v for k, v in fields_of(profile).items() if _mentioned(k, printed)}


def _mentioned(key: str, printed: str) -> bool:
    """這個項目可以列給模型挑嗎？名單裡沒提到的就不列。

    攤平後是 certificate[0].issued、欄位代碼是 certificate[].issued，名單寫的是後者：
    比對前先把序號抹掉，否則「只擋某一個欄位」那幾條永遠對不上，等於沒擋。
    """
    key = _ROW_INDEX.sub("[]", key)
    flat = _squash(printed)
    return not any(key.startswith(root) and not any(w in flat for w in words)
                   for root, words in ASK_ONLY_IF_MENTIONED.items())


def _worth_offering(key: str, fields: Dict[str, str]) -> bool:
    """一列一筆的表：名單裡單獨指名的那幾個欄位，還要真的填了值才列給模型挑。

    「發照日期」這一欄表格有問、但使用者沒填，列出來只會佔掉一欄——模型把那一欄
    認成發照日期，真正有資料的代碼就沒地方去了（實測履歷表因此少填兩格證照名稱）。
    整個區段的名單（家人、諮詢人）不套這一條：那問的是「這份表有沒有要這個人的資料」，
    跟填了沒有是兩回事。
    """
    if not any(key.startswith(root) and "." in root for root in ASK_ONLY_IF_MENTIONED):
        return True
    return any(v for k, v in fields.items() if _ROW_INDEX.sub("[]", k) == key)


class Cancelled(Exception):
    """使用者按了取消。分析在批次之間檢查，不會卡在模型那一個呼叫裡。"""


def analyze(blank: Path, profile: Dict[str, Any],
            host: str = LLM_HOST, model: str = LLM_MODEL,
            progress: Optional[Any] = None, stop: Optional[Any] = None) -> Draft:
    """看著版面決定每一個位置放哪一項資料。不寫檔——寫檔是 write() 的事，
    中間留給使用者修正。

    progress(說明文字)：每一批問完回報一次，讓使用者看得到「第幾批／共幾批」。
    stop()：回 True 就中止（使用者按了取消）。兩個都在批次之間呼叫，
    不會打斷正在跑的那一次模型呼叫。
    """
    def step(text: str) -> None:
        if stop is not None and stop():
            raise Cancelled(text)
        if progress is not None:
            progress(text)

    doc, form, all_slots = parse(blank)
    # 「這次應徵」才認出來的位置不進這一輪：模型看到的位置與批次跟以前一模一樣
    slots = [s for s in all_slots if not s.extra]
    fields = usable_fields(form, profile)
    pages = render_pages(doc)
    log.info("%s：可寫位置 %d 處、資料 %d 項、示意圖 %d 張",
             blank.name, len(slots), len(fields), len(pages))

    chosen: Dict[str, str] = {}
    ticks: Dict[str, bool] = {}
    # 一格一個字的那一排只問第一格，其餘跟著它寫
    followers = frozenset(x.id for run in char_runs(slots).values() for x in run[1:])
    # 一列一筆的表先整張問「每一欄是什麼」，那些格子就不再逐格問、也不參與補漏
    blocks = row_blocks(slots)
    block_cells = {x.addr for group in blocks.values() for x in group}
    in_blocks = frozenset(x.id for x in slots if x.addr in block_cells)
    if blocks:
        step("判讀一列一筆的表")
        try:
            rows, marks = ask_rows(blocks, fields, pages, printed_text(form), host, model)
            chosen.update(rows)
            ticks.update(marks)
        except llm.LlmError as e:
            log.warning("一列一筆的表判讀失敗，改回逐格問：%s", e)
            in_blocks = frozenset()
    # 勾選題自己問一輪。排在逐格問之後，才看得到同一列已經填了哪些資料
    groups = box_groups([x for x in slots if x.id not in in_blocks])
    in_boxes = frozenset(x.id for g in groups.values() for x in g)
    # 這幾批共用一串對話：提示快取只有「新提示完整包含舊提示」才重用，各開一份
    # 就每一批都要把示意圖重算一次（見 llm.Chat）
    pairing = llm.Chat(host, SYSTEM, model)
    batches = list(_batches([x for x in slots
                             if x.id not in in_blocks | in_boxes | followers]))
    for n, batch in enumerate(batches, 1):
        step(f"逐格判讀 第 {n}／{len(batches)} 批")
        try:
            chosen.update(ask(batch, fields, pages, chosen, pairing))
        except llm.LlmError as e:      # 一批失敗不該讓整份表格陪葬，其餘照跑、照評分
            log.warning("第 %d 批問失敗：%s", n, e)

    if groups:
        step(f"判讀勾選題（{len(groups)} 題）")
        try:
            picked, tk = ask_boxes(groups, fields, pages, chosen,
                                   host=host, model=model)
            ticks.update(tk)
            # 沒答出來的再問一次，而且只列那一題的候選。十幾題一起問時模型會跳過
            # 幾題（同一題換一批問又答得出來），選項縮到四個以內就是小得多的一道題
            again = {g: bs for g, bs in groups.items() if not any(b.id in picked for b in bs)}
            if again and len(again) < len(groups):
                more, tk = ask_boxes(again, fields, pages, {**chosen, **picked},
                                     narrow=True, host=host, model=model)
                picked.update(more)
                ticks.update(tk)
            picked.update(obvious_boxes(groups, fields, picked, ticks))
            one_question_per_field(groups, fields, picked, ticks)
            chosen.update(picked)
        except llm.LlmError as e:
            log.warning("勾選題判讀失敗：%s", e)
    skip = in_blocks | in_boxes | followers

    # 補漏看去重之後的結果：第一輪被丟掉的配對（住家電話欄填了已經用過的行動電話）
    # 讓出來的位置，也要能再問一次
    kept = dedupe(slots, chosen, fields, trusted=in_blocks)
    step("補漏")
    try:
        kept = dedupe(slots, {**kept, **recall(slots, kept, fields, pages, skip,
                                               host, model)},
                      fields, trusted=in_blocks)
    except llm.LlmError as e:
        log.warning("補漏那一輪失敗：%s", e)
    ticks = {k: v for k, v in ticks.items() if k in kept}   # 去重丟掉的配對，勾選也跟著不算
    # 應徵職務、工作地點那幾格由「這次應徵」面板決定，模型挑的不算——
    # 那些值不在「我的資料」裡，模型挑什麼都是別的欄位的資料
    per_job = {s.id for s in all_slots if s.job_field}
    kept = {sid: key for sid, key in kept.items() if sid not in per_job}
    ticks = {sid: tick for sid, tick in ticks.items() if sid not in per_job}
    log.info("%s：模型指定 %d 處、留下 %d 處", blank.name, len(chosen), len(kept))
    return Draft(slots=all_slots, fields=fields, assignment=kept, ticks=ticks)


def form_slots(slots: List[Slot]) -> List[Slot]:
    """算格式指紋用的位置：不含「這次應徵」才認出來的那些，
    否則同一份表格的指紋會因為這個功能而改掉，學過的格式全部要重學。"""
    return [s for s in slots if not s.extra]


def job_assignment(slots: List[Slot], values: Dict[str, str]) -> Dict[str, str]:
    """「這次應徵」那幾格要填什麼：面板填了的才寫，沒填的留白（跟以前一樣）。
    values 是攤平的 {"job.title": "資深工程師"}。"""
    return {s.id: s.job_field for s in slots
            if s.job_field and str(values.get(s.job_field, "")).strip()}


def write(blank: Path, out: Path, assignment: Dict[str, str],
          ticks: Dict[str, bool], profile: Dict[str, Any],
          highlight: bool = False, typed: Optional[Dict[str, str]] = None) -> int:
    """把決定好的值寫回文件，存到 out。回傳實際寫了幾處。

    重新解析一份原檔再寫，所以使用者改過 assignment 之後可以再叫一次；
    預覽與正式匯出也是各寫各的，不會互相汙染。
    """
    doc, form, slots = parse(blank)
    fields = usable_fields(form, profile)
    if typed:
        # 自己打的值沒有欄位代碼，給它一個；這樣底下的寫入完全不必分兩套
        fields = {**fields, **{TYPED + sid: text for sid, text in typed.items()}}
        assignment = {**assignment, **{sid: TYPED + sid for sid in typed}}
    written = apply_fills(slots, assignment, fields, ticks, highlight)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return written


def solve(blank: Path, profile: Dict[str, Any], out: Path,
          host: str = LLM_HOST, model: str = LLM_MODEL) -> Dict[str, Any]:
    """分析完直接寫檔。研究迴圈的評分走這個入口。"""
    draft = analyze(blank, profile, host, model)
    written = write(blank, out, draft.assignment, draft.ticks, profile)
    log.info("%s：實際寫入 %d 處", blank.name, written)
    # 配對結果一起回報：分數退步時要分得出是「模型配錯」還是「程式寫錯」
    return {"slots": len(draft.slots), "chosen": len(draft.assignment),
            "written": written, "assignment": draft.assignment, "ticks": draft.ticks}
