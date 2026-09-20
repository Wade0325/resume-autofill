"""大頭照：存一張，填履歷時自動貼進照片格。

照片是「加上去」的，印好的字一個都不能動——所以每次貼完都跑一次完整性檢查。
"""
from __future__ import annotations

import io
import shutil
import sys
from contextlib import redirect_stdout

import docx
import pytest
from PIL import Image

from backend import db, service
from backend.core import photo

sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parent.parent / "tools"))
import check_docx_integrity as integrity  # noqa: E402


def fake_photo(size=(900, 1200), fmt="PNG") -> bytes:
    """一張假的大頭照。畫點東西免得整張純色被當成空白圖。"""
    img = Image.new("RGB", size, (200, 190, 180))
    for x in range(0, size[0], 40):
        for y in range(0, size[1], 40):
            img.putpixel((x, y), (60, 60, 60))
    buf = io.BytesIO()
    img.save(buf, fmt)
    return buf.getvalue()


@pytest.fixture
def form_with_photo_cell(tmp_path):
    """一份有照片格的虛構表格。照片格就是印著「脫帽照片」那類字的儲存格。"""
    d = docx.Document()
    d.add_paragraph("虛構公司 應徵人員資料表")
    t = d.add_table(rows=2, cols=3)
    t.style = "Table Grid"
    t.cell(0, 0).text = "姓　名"
    t.cell(0, 1).text = ""
    t.cell(0, 2).text = "最近三個月內二吋脫帽照片"
    t.cell(1, 0).text = "行動電話"
    t.cell(1, 1).text = ""
    t.cell(0, 2).merge(t.cell(1, 2))        # 合併的照片格要算一格，不是兩格
    path = tmp_path / "with_photo.docx"
    d.save(str(path))
    return path


class TestSaving:
    def test_converts_to_jpeg_and_shrinks(self):
        service.save_photo(fake_photo((2400, 3200)))
        saved = service.photo_path()
        assert saved.name == "photo.jpg"
        with Image.open(saved) as img:      # Windows 上開著檔案就刪不掉，檢查完馬上關
            assert img.format == "JPEG"
            assert max(img.size) <= 1200
            # 手機拍的照片帶著機型與拍攝地點，履歷是要寄出去的
            assert dict(img.getexif()) == {}

    def test_rejects_non_images(self):
        with pytest.raises(ValueError, match="圖片"):
            service.save_photo("這不是圖片".encode("utf-8"))


class TestPhotoCell:
    def test_finds_the_cell(self, form_with_photo_cell):
        cells = photo.photo_cells(docx.Document(str(form_with_photo_cell)))
        assert len(cells) == 1              # 合併的算一格
        assert "脫帽照片" in cells[0].text

    def test_inserts_and_scales(self, form_with_photo_cell, tmp_path):
        service.save_photo(fake_photo())
        out = tmp_path / "filled.docx"
        shutil.copy(form_with_photo_cell, out)

        assert photo.insert(out, service.photo_path()) == 1
        after = docx.Document(str(out))
        assert len(after.inline_shapes) == 1
        assert round(after.inline_shapes[0].width.cm, 1) <= 3.5   # 二吋照片的寬度

    def test_printed_text_survives(self, form_with_photo_cell, tmp_path):
        service.save_photo(fake_photo())
        out = tmp_path / "filled.docx"
        shutil.copy(form_with_photo_cell, out)
        photo.insert(out, service.photo_path())

        after = docx.Document(str(out))
        text = "\n".join(c.text for t in after.tables for r in t.rows for c in r.cells)
        assert "脫帽照片" in text

    def test_nothing_else_is_damaged(self, form_with_photo_cell, tmp_path):
        """完整性檢查：勾選框、底線、字型、段落數都要跟原稿一樣。"""
        service.save_photo(fake_photo())
        out = tmp_path / "filled.docx"
        shutil.copy(form_with_photo_cell, out)
        photo.insert(out, service.photo_path())

        report = io.StringIO()
        with redirect_stdout(report):
            intact = integrity.compare(form_with_photo_cell, out)
        assert intact, report.getvalue()


class TestWholeFlow:
    def test_output_carries_the_photo(self, client, make_job):
        service.save_photo(fake_photo())
        job_id = make_job(decided={})
        client.post(f"/api/jobs/{job_id}/output", json={})
        # 這份樣本表格沒有照片格，所以成品不該有圖——沒照片格就不亂貼
        assert len(docx.Document(str(service.output_path(job_id))).inline_shapes) == 0

    def test_backup_round_trip(self, client):
        service.save_photo(fake_photo())
        body = client.get("/api/profile/export").json()
        assert body.get("photo")

        service.delete_photo()
        assert service.photo_path().exists() is False

        assert client.post("/api/profile/restore", json=body).status_code == 200
        assert service.photo_path().exists()


class TestPhotoApi:
    def test_upload_read_delete(self, client):
        assert client.post("/api/profile/photo",
                           files={"file": ("me.png", fake_photo(), "image/png")}).status_code == 200
        assert client.get("/api/profile/photo").headers["content-type"] == "image/jpeg"
        assert client.delete("/api/profile/photo").status_code == 200
        assert client.get("/api/profile/photo").status_code == 404
        assert client.delete("/api/profile/photo").status_code == 404

    def test_rejects_bad_extension(self, client):
        r = client.post("/api/profile/photo", files={"file": ("x.txt", b"hello", "text/plain")})
        assert r.status_code == 400

    def test_rejects_fake_image(self, client):
        """副檔名對但內容不是圖片——只看副檔名是不夠的。"""
        r = client.post("/api/profile/photo",
                        files={"file": ("x.png", b"not an image", "image/png")})
        assert r.status_code == 422

    def test_profile_is_untouched(self, client):
        from tests.conftest import PROFILE

        client.post("/api/profile/photo",
                    files={"file": ("me.png", fake_photo(), "image/png")})
        assert db.get_kv("profile") == PROFILE
