"""學過的格式：怎麼認、列得出來、刪得掉、可以不用它重新判讀。

重點是指紋要含欄名——只看結構的話，同一套版型的不同公司會共用同一份對映，填出來全錯。
"""
from __future__ import annotations

import time

import docx
import pytest

from backend import db, service
from backend.core import document, filler

MAPPING = {"field_key": "basic.name_zh", "ordinal": 0, "label": "姓名"}


def make_form(path, labels):
    """同一套結構、只有欄名不同的表格。"""
    d = docx.Document()
    d.add_paragraph("虛構公司 應徵資料表")
    t = d.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    for i, label in enumerate(labels):
        t.cell(i, 0).text = label
        t.cell(i, 1).text = ""
    d.save(str(path))
    return path


def keys_of(path):
    """(只看結構的舊指紋, 結構＋欄名的新鍵)"""
    _d, form, slots = filler.parse(path)
    structure = "vlm:" + document.fingerprint(filler.form_slots(slots))
    return structure, service._template_key(structure, filler.printed_text(form))


@pytest.fixture
def two_forms(tmp_path):
    return (make_form(tmp_path / "a.docx", ["姓名", "電話"]),
            make_form(tmp_path / "b.docx", ["職稱", "分機"]))


class TestFingerprint:
    def test_same_layout_different_labels_no_longer_share(self, two_forms):
        a, b = two_forms
        (structure_a, key_a), (structure_b, key_b) = keys_of(a), keys_of(b)
        # 結構一樣——以前就是靠這個當鍵，所以兩家公司會共用同一份對映
        assert structure_a == structure_b
        assert key_a != key_b

    def test_same_document_is_stable(self, two_forms):
        a, _b = two_forms
        assert keys_of(a)[1] == keys_of(a)[1]

    def test_legacy_key_is_still_found(self, two_forms):
        """舊資料只有結構指紋，不能因為改了鍵就全部作廢。"""
        a, b = two_forms
        structure_a, key_a = keys_of(a)
        structure_b, key_b = keys_of(b)
        db.put_template(structure_a, {"t0.r0.c1#0.1": MAPPING}, source_name="舊版學的.docx")
        assert service._cached_template(key_a, structure_a)
        # 別份表格不該撈到別人的——這正是當初的 bug
        assert service._cached_template(key_b, structure_b) == {}


class TestManageApi:
    @pytest.fixture(autouse=True)
    def _two_templates(self, two_forms):
        # 兩份都存「結構＋欄名」的新鍵。刻意不留舊鍵的那一筆——這兩份表格結構相同，
        # 留著的話忘掉其中一份之後，舊鍵回退會把它又撈回來（那是另一項測試在測的事）
        a, b = two_forms
        db.put_template(keys_of(a)[1], {"t0.r0.c1#0.1": MAPPING}, source_name="舊版學的.docx")
        db.put_template(keys_of(b)[1], {"t0.r0.c1#0.1": MAPPING}, source_name="新版學的.docx")
        self.key_b = keys_of(b)[1]
        self.structure_b = keys_of(b)[0]

    def test_lists_both(self, client):
        rows = client.get("/api/templates").json()
        assert sorted(r["source_name"] for r in rows) == ["新版學的.docx", "舊版學的.docx"]
        assert {r["engine"] for r in rows} == {"vlm"}
        assert {r["slots"] for r in rows} == {1}

    def test_forget_removes_it(self, client):
        assert client.post("/api/templates/forget", json={"fingerprint": self.key_b}).status_code == 200
        assert service._cached_template(self.key_b, self.structure_b) == {}
        assert len(client.get("/api/templates").json()) == 1

    def test_forget_unknown_is_404(self, client):
        r = client.post("/api/templates/forget", json={"fingerprint": "沒有這個"})
        assert r.status_code == 404


class TestReanalyze:
    def test_learned_format_fills_without_a_model(self, client, two_forms):
        """模型沒開，但這份格式學過——照樣分析得出來。"""
        a, _b = two_forms
        job_id = self._job_from(a)
        service._vlm_worker(job_id, "a.docx")

        state = client.get(f"/api/jobs/{job_id}").json()
        assert state["status"] == "ready"
        assert state["plan"]["template_cached"]

    def test_reanalyze_skips_the_cache(self, client, two_forms):
        """要求重新判讀就真的去問模型——模型沒開，所以這次會失敗。"""
        a, _b = two_forms
        job_id = self._job_from(a)
        service._vlm_worker(job_id, "a.docx")

        assert client.post(f"/api/jobs/{job_id}/reanalyze").status_code == 200
        state = self._wait(client, job_id)
        assert state["status"] == "failed"
        assert "模型還沒啟動" in state.get("error", "")
        # 重新判讀不該把學過的格式刪掉
        assert service._cached_template(*reversed(keys_of(a)))

    def test_reanalyze_unknown_is_404(self, client):
        assert client.post("/api/jobs/aaaaaaaaaaaa/reanalyze").status_code == 404

    @staticmethod
    def _job_from(form_path):
        import shutil
        import uuid

        job_id = uuid.uuid4().hex[:12]
        service.job_dir(job_id).mkdir(parents=True, exist_ok=True)
        shutil.copy(form_path, service.input_path(job_id))
        _d, _f, slots = filler.parse(form_path)
        db.put_template(keys_of(form_path)[1],
                        {s.id: MAPPING for s in slots}, source_name="a.docx")
        db.create_job(job_id, "a.docx", status="processing", engine="vlm")
        return job_id

    @staticmethod
    def _wait(client, job_id, tries=40):
        for _ in range(tries):
            state = client.get(f"/api/jobs/{job_id}").json()
            if state["status"] != "processing":
                return state
            time.sleep(0.25)
        return state
