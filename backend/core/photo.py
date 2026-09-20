# -*- coding: utf-8 -*-
"""把大頭照貼進履歷的照片格。

表格上的照片格認得出來：格子裡印著「照片」「相片」「脫帽照」這類字。
照片是「加上去」的，印好的字一個都不動——跟寫字那條路一樣的原則。
"""
from __future__ import annotations

import copy
import logging
import re
from pathlib import Path
from typing import Any, List

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Emu

log = logging.getLogger(__name__)

# 照片格印的字。「照」單獨一個字太容易誤中（「對照」「按照」），一律看兩個字以上
PHOTO_RE = re.compile(r"照片|相片|照\s*片|脫帽照|大頭照|近照")
# 二吋照片約 3.5×4.5 cm；格子更窄就跟著縮，更寬也不要放大到誇張
MAX_WIDTH_CM = 3.5
MARGIN_CM = 0.3          # 照片與格線之間留一點，不要貼著框
FALLBACK_WIDTH_CM = 3.0  # 格子沒寫寬度時（tcW 缺）用這個


def photo_cells(doc: Any) -> List[Any]:
    """這份文件裡哪些格子是照片格。合併儲存格會在好幾個座標上重複出現，只算一次。"""
    out, seen = [], set()
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                if PHOTO_RE.search(cell.text):
                    out.append(cell)
    return out


def _width_of(cell: Any) -> Emu:
    """照片要多寬：格子扣掉邊距，最多 3.5 公分（二吋照片的寬度）。"""
    usable = Cm(FALLBACK_WIDTH_CM)
    if cell.width:
        usable = Emu(max(int(cell.width) - int(Cm(MARGIN_CM)), int(Cm(1))))
    return Emu(min(int(usable), int(Cm(MAX_WIDTH_CM))))


def insert(path: Path, photo: Path) -> int:
    """把 photo 貼進 path 這份文件的照片格，存回原檔。回傳貼了幾格。

    照片接在格子最後一段的後面（換行再放），印好的說明（「最近半年內二吋半身脫帽照片」）
    留著不動。不另外開一段：完整性檢查是拿段落一段一段對的，多一段就整份對不上了。
    新的 run 沿用同段既有的字型設定，免得被當成「沒有字型的 run」。
    """
    doc = Document(str(path))
    cells = photo_cells(doc)
    for cell in cells:
        para = cell.paragraphs[-1]
        run = para.add_run()
        donor = next((r for r in para.runs if r._r.find(qn("w:rPr")) is not None), None)
        if donor is not None:
            run._r.insert(0, copy.deepcopy(donor._r.find(qn("w:rPr"))))
        if para.text.strip():
            run.add_break()        # 說明文字底下才是照片
        run.add_picture(str(photo), width=_width_of(cell))
    if cells:
        doc.save(str(path))
        log.info("貼上照片 %d 格", len(cells))
    return len(cells)
